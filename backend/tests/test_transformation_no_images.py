"""
Material sin imágenes (Issue #117): ubicación de cada descripción en el texto, imágenes decorativas,
información solo por color, material sin imágenes, textos alternativos del diagnóstico y exportación.
La IA con visión se reemplaza por respuestas fijas.
"""
import asyncio
import io
import time
from unittest.mock import AsyncMock, patch
from urllib.parse import urlsplit

import pymupdf
import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
from fastapi.testclient import TestClient

from app.core import rate_limit
from app.core.config import settings
from app.db.database import init_db
from app.main import app
from app.schemas.transformation import ImageDescription, NoImagesAIOutput
from app.services import storage_service
from app.services.llm.base import Usage
from app.services.llm_service import StructuredResult
from app.services.storage_service import LocalStorage
from app.services.transformations import no_images
from app.services.transformations.base import SourceMaterial

DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
INTRO = "El gráfico muestra la cantidad de lluvia caída en cada estación del año en la provincia."
MIDDLE = "Compará los valores de cada estación y anotá en qué época llueve más."
CLOSING = "Respondé las preguntas del final con tus propias palabras."


def _png(width: int = 300, height: int = 200) -> bytes:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, width, height), False)
    pix.set_rect(pix.irect, (40, 120, 200))
    return pix.tobytes("png")


def _docx_with_images() -> bytes:
    document = Document()
    document.add_heading("Las lluvias en la provincia", level=1)
    document.add_paragraph(INTRO)
    document.add_picture(io.BytesIO(_png()), width=Inches(3))
    document.add_paragraph(MIDDLE)
    document.add_picture(io.BytesIO(_png(320, 210)), width=Inches(2))
    document.add_paragraph(CLOSING)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


AI_OUTPUT = {
    "images": [
        {
            "image_index": 1, "role": "informativa",
            "description": "Gráfico de barras con la lluvia de cada estación: el verano es la estación más lluviosa.",
            "data_table": "| Estación | Lluvia (mm) |\n| --- | --- |\n| Verano | 320 |\n| Invierno | 40 |",
            "color_only": "Las barras en azul oscuro marcan las estaciones con más lluvia que el promedio.",
            "insert_after": "en cada estación del año en la provincia",
        },
        {
            "image_index": 2, "role": "decorativa", "description": "Dibujo de una nube",
            "data_table": "", "color_only": "", "insert_after": "",
        },
    ]
}


def _ai_result(output: dict = AI_OUTPUT) -> StructuredResult:
    return StructuredResult(
        parsed=NoImagesAIOutput.model_validate(output), provider="gemini", model="gemini-flash-test",
        connection_label="Gemini", latency_ms=900, usage=Usage(prompt_tokens=1500, output_tokens=200, total_tokens=1700),
    )


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "DATABASE_PATH", str(tmp_path / "test.db"))
    monkeypatch.setattr(settings, "STORAGE_BACKEND", "local")
    monkeypatch.setattr(settings, "FILE_CLEANUP_INTERVAL_MINUTES", 0)
    init_db()
    storage_service.reset_storage(LocalStorage(str(tmp_path / "storage"), b"secreto-de-prueba"))
    rate_limit.reset()
    with TestClient(app) as c:
        yield c
    storage_service.reset_storage()
    rate_limit.reset()


def _upload(client, content: bytes, filename: str = "lluvias.docx", content_type: str = DOCX_TYPE) -> str:
    data = client.post("/api/v1/files/upload-url",
                       json={"filename": filename, "content_type": content_type, "size_bytes": len(content)}).json()
    url = urlsplit(data["upload"]["url"])
    assert client.put(f"{url.path}?{url.query}", content=content, headers=data["upload"]["headers"]).status_code == 200
    assert client.post(f"/api/v1/files/{data['file']['file_id']}/complete").status_code == 200
    return data["file"]["file_id"]


def _wait(client, transformation_id: str) -> dict:
    deadline = time.monotonic() + 5
    while time.monotonic() < deadline:
        data = client.get(f"/api/v1/transformations/{transformation_id}").json()
        if data["status"] in ("ready", "failed"):
            return data
        time.sleep(0.02)
    raise AssertionError("La transformación no terminó a tiempo")


def _transform(client, *, mock: bool = False, output: dict = AI_OUTPUT, **source):
    with patch("app.services.llm_service.llm_service.generate_structured", AsyncMock(return_value=_ai_result(output))) as llm:
        res = client.post(f"/api/v1/transformations?mock={str(mock).lower()}", json={"format": "no_images", **source})
        assert res.status_code == 202, res.text
        return _wait(client, res.json()["transformation_id"]), llm


def _image(index: int, role: str = "informativa", **fields) -> ImageDescription:
    return ImageDescription(image_index=index, location=f"Página 1, imagen {index}", role=role,
                            description=fields.pop("description", f"Descripción {index}"), source="ia", **fields)


# --- Ubicación de las descripciones ---------------------------------------------------------


TEXT = f"Título\n{INTRO}\n{MIDDLE}\n{CLOSING}"


def test_description_goes_after_the_quoted_paragraph():
    result, _ = no_images.build(TEXT, [(_image(1), ["«Estación del año en la PROVINCIA.»"])])
    lines = result.text.splitlines()
    assert lines[1] == INTRO
    assert lines[3] == "[Imagen 1 · Página 1, imagen 1]"
    assert lines[4] == "Descripción 1"
    assert lines[5] == "[Fin de la imagen 1]"
    assert lines[7] == MIDDLE
    assert result.images[0].placement == "en_su_lugar"


def test_falls_back_to_the_next_anchor_and_then_to_the_end():
    items = [
        (_image(1), ["una cita que no está en el material", "anotá en qué época llueve más"]),
        (_image(2), ["tampoco está", ""]),
    ]
    result, warnings = no_images.build(TEXT, items)
    text = result.text
    assert text.index(MIDDLE) < text.index("[Imagen 1") < text.index(CLOSING)
    assert text.index(CLOSING) < text.index(no_images._UNPLACED_TITLE) < text.index("[Imagen 2")
    assert [i.placement for i in result.images] == ["en_su_lugar", "al_final"]
    assert "no se pudo ubicar" in warnings[0]


def test_repeated_instructions_keep_each_image_in_its_activity():
    instruction = "Observá el gráfico y respondé las preguntas."
    text = f"Actividad 1\n{instruction}\nPregunta de la actividad 1.\nActividad 2\n{instruction}\nPregunta de la actividad 2."
    items = [(_image(1), [instruction]), (_image(2), [instruction])]
    result, warnings = no_images.build(text, items)
    out = result.text
    assert out.index("Actividad 1") < out.index("[Imagen 1") < out.index("Actividad 2") < out.index("[Imagen 2")
    assert [i.placement for i in result.images] == ["en_su_lugar", "en_su_lugar"]
    assert not warnings


def test_ellipsis_quote_searches_each_part_after_the_previous_one():
    text = f"Título\nEn la provincia llueve más en verano.\n{INTRO}\n{MIDDLE}"
    # "en la provincia" ya aparece en el segundo renglón: tiene que buscarse después de "El gráfico muestra…"
    result, _ = no_images.build(text, [(_image(1), ["El gráfico muestra la cantidad... en la provincia"])])
    lines = result.text.splitlines()
    assert lines.index(INTRO) < lines.index("[Imagen 1 · Página 1, imagen 1]") < lines.index(MIDDLE)


def test_two_images_after_the_same_paragraph_stay_together():
    result, warnings = no_images.build(TEXT, [(_image(1), [INTRO]), (_image(2), [INTRO])])
    out = result.text
    assert out.index(INTRO) < out.index("[Imagen 1") < out.index("[Imagen 2") < out.index(MIDDLE)
    assert not warnings


def test_short_anchors_are_ignored():
    result, _ = no_images.build(TEXT, [(_image(1), ["Título"])])
    assert result.images[0].placement == "al_final"


def test_decorative_images_are_omitted_and_reported():
    result, warnings = no_images.build(TEXT, [(_image(1, "decorativa", description="Dibujo de una nube"), [INTRO])])
    assert result.text == TEXT
    assert result.omitted_images == ["Página 1, imagen 1: Dibujo de una nube"]
    assert result.images[0].placement is None
    assert "decorativa" in warnings[0]


def test_data_table_and_color_only_go_in_the_text():
    image = _image(1, data_table="| A | B |\n| --- | --- |\n| 1 | 2 |", color_only="En rojo, las zonas secas.")
    result, warnings = no_images.build(TEXT, [(image, [INTRO])])
    assert "Datos:\n| A | B |" in result.text
    assert "Información que la imagen daba solo con color: En rojo, las zonas secas." in result.text
    assert result.color_notes == ["Imagen 1 (Página 1, imagen 1): En rojo, las zonas secas."]
    assert any("solo con color" in w for w in warnings)


def test_context_anchor_takes_the_text_before_the_image():
    assert no_images._context_anchor("uno dos tres [...] después de la imagen") == "uno dos tres"
    assert no_images._context_anchor("a b c d e f g h i j") == "c d e f g h i j"
    assert no_images._context_anchor("palabras completas y una cort...") == "palabras completas y una"


def test_material_without_images_stays_equal_and_skips_the_ai():
    with patch("app.services.llm_service.llm_service.generate_structured", AsyncMock()) as llm:
        generation = asyncio.run(no_images.generate(SourceMaterial(text=TEXT)))
    llm.assert_not_awaited()
    assert generation.result.text == TEXT
    assert generation.result.images == []
    assert generation.metadata.total_tokens == 0
    assert "igual al original" in generation.warnings[0]


# --- Endpoint ---------------------------------------------------------------------------


def test_file_images_are_described_with_vision_in_place(client):
    file_id = _upload(client, _docx_with_images())
    data, llm = _transform(client, file_id=file_id)
    assert data["status"] == "ready", data.get("error")

    kwargs = llm.await_args.kwargs
    assert kwargs["require_vision"] is True
    assert kwargs["schema"] is NoImagesAIOutput
    assert sum(1 for part in kwargs["contents"] if part["type"] == "image") == 2

    result = data["result"]
    text = result["text"]
    assert text.index(INTRO) < text.index("[Imagen 1") < text.index(MIDDLE)
    assert "el verano es la estación más lluviosa" in text
    assert "| Verano | 320 |" in text
    assert "Dibujo de una nube" not in text
    assert result["omitted_images"] == [f"{result['images'][1]['location']}: Dibujo de una nube"]
    assert result["images"][0]["source"] == "ia"
    assert result["color_notes"]
    assert data["metadata"]["model"] == "gemini-flash-test"


def test_image_missing_from_the_ai_answer_is_flagged(client):
    file_id = _upload(client, _docx_with_images())
    data, _ = _transform(client, output={"images": AI_OUTPUT["images"][:1]}, file_id=file_id)
    assert data["status"] == "ready"
    assert any("no describió la imagen 2" in w for w in data["warnings"])
    assert "Falta la descripción de esta imagen" in data["result"]["text"]


def test_analysis_uses_the_diagnosis_alt_texts_without_ai(client):
    file_id = _upload(client, _docx_with_images())
    analysis = client.post(f"/api/v1/analysis/files/{file_id}?mock=true").json()
    assert analysis["alt_text_proposals"]

    data, llm = _transform(client, analysis_id=analysis["analysis_id"])
    llm.assert_not_awaited()
    assert data["status"] == "ready"
    images = data["result"]["images"]
    assert [i["source"] for i in images] == ["diagnostico"] * len(analysis["alt_text_proposals"])
    assert analysis["alt_text_proposals"][0]["suggested_alt_text"] in data["result"]["text"]
    assert any("textos alternativos que propuso el diagnóstico" in w for w in data["warnings"])


def test_mock_places_simulated_descriptions(client):
    data, llm = _transform(client, mock=True, file_id=_upload(client, _docx_with_images()))
    llm.assert_not_awaited()
    text = data["result"]["text"]
    assert "[MOCK] Descripción simulada de la imagen 1." in text
    assert text.index(INTRO) < text.index("[Imagen 1")


@pytest.mark.parametrize("fmt,content_type,signature", [
    ("docx", DOCX_TYPE, b"PK"), ("pdf", "application/pdf", b"%PDF-"),
    ("md", "text/markdown", b"# Material sin im"), ("txt", "text/plain", b"MATERIAL SIN IM"),
])
def test_accepted_version_exports_like_the_adapted_material(client, fmt, content_type, signature):
    data, _ = _transform(client, file_id=_upload(client, _docx_with_images()))
    tid = data["transformation_id"]
    client.post(f"/api/v1/transformations/{tid}/decision", json={"decision": "aceptada"})
    res = client.get(f"/api/v1/transformations/{tid}/export?format={fmt}")
    assert res.status_code == 200
    assert res.headers["content-type"].startswith(content_type)
    assert res.content.startswith(signature)
    if fmt in ("md", "txt"):
        assert "el verano es la estación más lluviosa" in res.text
        assert "Imagen decorativa omitida" in res.text


def test_data_table_is_a_real_table_in_docx(client):
    data, _ = _transform(client, file_id=_upload(client, _docx_with_images()))
    tid = data["transformation_id"]
    client.post(f"/api/v1/transformations/{tid}/decision", json={"decision": "aceptada"})
    res = client.get(f"/api/v1/transformations/{tid}/export?format=docx")
    tables = Document(io.BytesIO(res.content)).tables
    assert len(tables) == 1
    assert [[c.text for c in row.cells] for row in tables[0].rows] == [
        ["Estación", "Lluvia (mm)"], ["Verano", "320"], ["Invierno", "40"],
    ]
    assert tables[0].rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None


def test_file_warnings_skip_the_scanner_notice(client):
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), INTRO, fontsize=9)
    page.insert_image(pymupdf.Rect(72, 100, 372, 300), stream=_png())
    page.insert_text((72, 330), CLOSING, fontsize=9)
    file_id = _upload(client, pdf.tobytes(), filename="lluvias.pdf", content_type="application/pdf")
    data, _ = _transform(client, file_id=file_id)
    assert data["status"] == "ready", data.get("error")
    assert not any("escáner de accesibilidad" in w for w in data["warnings"])


def test_material_without_ai_call_does_not_use_the_ai_quota(client, monkeypatch):
    monkeypatch.setattr(settings, "AI_RATE_LIMIT_PER_IP_DAY", 1)
    analysis_id = client.post("/api/v1/analysis/sample?mock=true", json={"text": TEXT}).json()["analysis_id"]
    for _ in range(2):  # sin imágenes no llama a la IA: el límite de IA no cuenta
        data, llm = _transform(client, analysis_id=analysis_id)
        llm.assert_not_awaited()
        assert data["status"] == "ready"
