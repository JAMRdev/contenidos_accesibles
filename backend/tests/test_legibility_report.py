"""
Revisión del PR #157 (issue #79):
- el informe final conserva las reglas de legibilidad del archivo (no hay "mejoras" sin cambios);
- el tope por barrera grave se conserva según ScorePenalty.severe, no según los puntos;
- una fuente de PDF que es solo el prefijo de subconjunto no cuenta como fuente.
"""

import io

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from fastapi.testclient import TestClient

from app.db.database import init_db
from app.main import app
from app.schemas.analysis import ScoreBreakdown, ScorePenalty
from app.services import rules_service as rs
from app.services.adaptation_service import _FILE_RULES, _score_after
from app.services.legibility_extractor import Fragment
from app.services.legibility_typography import RULE_ORDER, TypographyCollector
from app.services.office_service import office_service

TEXT = (
    "Los ecosistemas acuáticos reúnen seres vivos que dependen del agua para vivir. "
    "En esta actividad vamos a observar cómo se relacionan entre sí y con su ambiente. "
)
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@pytest.fixture(autouse=True)
def ensure_db():
    init_db()


def _docx_bytes(build) -> bytes:
    document = Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _with_every_legibility_problem(document):
    document.styles["Normal"].font.size = Pt(7)
    for _ in range(3):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = paragraph.add_run(TEXT * 2)
        run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
        run.italic = True


# --- Informe final -------------------------------------------------------------


def test_informe_sin_cambios_conserva_las_reglas_de_legibilidad():
    client = TestClient(app)
    analysis = client.post(
        "/api/v1/analysis/file?mock=true",
        files={"file": ("material.docx", _docx_bytes(_with_every_legibility_problem), DOCX)},
    ).json()
    report = client.get("/api/v1/adaptations/report", params={"analysis_id": analysis["analysis_id"]}).json()

    before = {p["rule"] for p in report["score_breakdown_before"]["penalties"]}
    after = {p["rule"] for p in report["score_breakdown_after"]["penalties"]}
    assert {"letra_chica_cuerpo", "contraste_insuficiente", "tipografia_cursiva"} <= before
    assert before <= after  # nada del archivo se pierde
    assert report["score_after"] == report["score_before"]
    assert report["score_delta"] == 0


def test_todas_las_reglas_de_legibilidad_son_reglas_del_archivo():
    # Si se agrega una regla de legibilidad nueva y no se suma a _FILE_RULES, este test falla.
    legibility_rules = {"letra_chica_cuerpo", "letra_chica_pasajes", "contraste_insuficiente"}
    legibility_rules |= {f"tipografia_{rule}" for rule in RULE_ORDER}
    assert legibility_rules <= _FILE_RULES


def _breakdown(*penalties: ScorePenalty) -> ScoreBreakdown:
    rules_score = max(0, 100 - sum(p.points for p in penalties))
    return ScoreBreakdown(
        ai_score=100, rules_score=rules_score, ai_weight=rs.AI_WEIGHT, rules_weight=rs.RULES_WEIGHT,
        penalties=list(penalties), final_score=0,
    )


def test_el_tope_por_contraste_grave_se_conserva_en_el_informe():
    severe = ScorePenalty(rule="contraste_insuficiente", category="accesibilidad_visual", count=1, points=8,
                          severe=True)
    after = _score_after(_breakdown(severe), TEXT)
    assert after.cap == rs.SEVERE_BARRIER_CAP
    assert after.final_score == rs.SEVERE_BARRIER_CAP


def test_dos_contrastes_leves_no_aplican_el_tope_aunque_sumen_8_puntos():
    # Dos pasajes leves (4 + 4) suman lo mismo que uno grave: los puntos no sirven para decidir el tope.
    mild = ScorePenalty(rule="contraste_insuficiente", category="accesibilidad_visual", count=2, points=8)
    after = _score_after(_breakdown(mild), TEXT)
    assert after.cap is None
    assert after.final_score > rs.SEVERE_BARRIER_CAP


def test_letra_chica_en_el_cuerpo_siempre_aplica_el_tope():
    # Diagnósticos guardados antes de ScorePenalty.severe: letra_chica_cuerpo está en _SEVERE_FILE_RULES.
    old = ScorePenalty.model_validate(
        {"rule": "letra_chica_cuerpo", "category": "accesibilidad_visual", "count": 1, "points": 15}
    )
    assert old.severe is False
    assert _score_after(_breakdown(old), TEXT).final_score == rs.SEVERE_BARRIER_CAP


def test_las_reglas_marcan_severe_solo_con_barreras_graves():
    def penalties(build) -> dict[str, bool]:
        extracted = office_service.process_docx(_docx_bytes(build), "material.docx")
        return {p.rule: p.severe for p in rs.run_rules(extracted).penalties}

    def mild_contrast(document):
        for _ in range(2):
            document.add_paragraph().add_run(TEXT).font.color.rgb = RGBColor(0x8A, 0x8A, 0x8A)  # 3,5:1

    def severe_contrast(document):
        document.add_paragraph().add_run(TEXT).font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)  # 1,6:1

    def decorative_body(document):
        document.styles["Normal"].font.name = "Lucida Handwriting"
        document.add_paragraph(TEXT * 2)

    assert penalties(mild_contrast)["contraste_insuficiente"] is False
    assert penalties(severe_contrast)["contraste_insuficiente"] is True
    assert penalties(decorative_body)["tipografia_fuente_decorativa"] is True
    assert penalties(_with_every_legibility_problem)["letra_chica_cuerpo"] is True


# --- Fuente que es solo el prefijo de subconjunto -------------------------------


def test_fuente_que_es_solo_el_prefijo_de_subconjunto_no_cuenta():
    collector = TypographyCollector()
    collector.add_paragraph([Fragment(text=TEXT * 3, font="ABCDEF+"), Fragment(text=TEXT, font="Calibri")], 1, None)
    body_font, decorative, _ = collector.result()
    assert body_font == "Calibri"  # antes ganaba None, por tener más caracteres
    assert decorative is False

    only_prefix = TypographyCollector()
    only_prefix.add_paragraph([Fragment(text=TEXT, font="ABCDEF+")], 1, None)
    assert only_prefix.result()[0] is None
