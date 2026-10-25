"""Contraste texto/fondo, WCAG 1.4.3 (issue #79, parte 2). Los archivos se generan en el test."""

import io

import pymupdf
import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGB
from pptx.util import Inches
from pptx.util import Pt as PptPt

from app.schemas.document import ExtractedDocument, LegibilityStats, LowContrastPassage
from app.services import rules_service as rs
from app.services.document_service import document_service
from app.services.legibility_extractor import contrast_ratio, is_large_text
from app.services.legibility_rules import MAX_CONTRAST_PENALTY, SEVERE_CONTRAST_POINTS
from app.services.office_service import office_service

TEXT = (
    "Los ecosistemas acuáticos reúnen seres vivos que dependen del agua para vivir. "
    "En esta actividad vamos a observar cómo se relacionan entre sí y con su ambiente. "
)
BLACK, WHITE, YELLOW, NAVY = (0, 0, 0), (255, 255, 255), (255, 255, 0), (0, 0, 128)
LIGHT_GRAY = (0xCC, 0xCC, 0xCC)  # 1,6:1 sobre blanco
MID_GRAY = (0x8A, 0x8A, 0x8A)  # 3,5:1 sobre blanco: falla en texto normal, pasa en texto grande
DARK_GRAY = (0x76, 0x76, 0x76)  # 4,5:1 sobre blanco: pasa


def _contrast(result: rs.RulesResult) -> list[dict]:
    return [f.barrier for f in result.findings if f.barrier["issue"].startswith("Contraste insuficiente")]


def _penalty(result: rs.RulesResult):
    return next((p for p in result.penalties if p.rule == "contraste_insuficiente"), None)


def _visual(result: rs.RulesResult):
    return next(d for d in result.dimensions if d.category == "accesibilidad_visual")


# --- Generadores de archivos -----------------------------------------------


def _pdf_bytes(*boxes) -> bytes:
    """boxes: (texto, color RGB 0-255, tamaño, fondo RGB o None)."""
    doc = pymupdf.open()
    page = doc.new_page()
    y = 40
    for text, color, size, background in boxes:
        rect = pymupdf.Rect(50, y, 550, y + 150)
        if background is not None:
            page.draw_rect(rect + (-5, -5, 5, 5), color=None, fill=tuple(v / 255 for v in background))
        page.insert_textbox(rect, text, fontsize=size, color=tuple(v / 255 for v in color))
        y += 170
    data = doc.tobytes()
    doc.close()
    return data


def _pdf(*boxes) -> rs.RulesResult:
    return rs.run_rules(document_service.process_pdf(_pdf_bytes(*boxes), "material.pdf"))


def _shd(fill: str):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    return shd


def _docx(*paragraphs, cell=None) -> rs.RulesResult:
    """paragraphs: (texto, color RGB o None, {'highlight': ..., 'shading': 'RRGGBB'})."""
    document = Document()
    for text, color, options in paragraphs:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(text)
        if color is not None:
            run.font.color.rgb = RGBColor(*color)
        if options.get("highlight"):
            run.font.highlight_color = options["highlight"]
        if options.get("shading"):
            paragraph._p.get_or_add_pPr().append(_shd(options["shading"]))
    if cell is not None:
        text, color, fill = cell
        table_cell = document.add_table(rows=1, cols=1).cell(0, 0)
        table_cell._tc.get_or_add_tcPr().append(_shd(fill))
        run = table_cell.paragraphs[0].add_run(text)
        run.font.color.rgb = RGBColor(*color)
    buffer = io.BytesIO()
    document.save(buffer)
    return rs.run_rules(office_service.process_docx(buffer.getvalue(), "material.docx"))


def _pptx_bytes(*boxes, slide_background=None) -> bytes:
    """boxes: (texto, color RGB o None, tamaño o None, relleno de la forma RGB o None)."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])  # en blanco
    if slide_background is not None:
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = PptRGB(*slide_background)
    top = 0.3
    for text, color, size, fill in boxes:
        shape = slide.shapes.add_textbox(Inches(0.5), Inches(top), Inches(9), Inches(2))
        if fill is not None:
            shape.fill.solid()
            shape.fill.fore_color.rgb = PptRGB(*fill)
        run = shape.text_frame.paragraphs[0].add_run()
        run.text = text
        if color is not None:
            run.font.color.rgb = PptRGB(*color)
        if size is not None:
            run.font.size = PptPt(size)
        top += 2.3
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _pptx(*boxes, **kwargs) -> rs.RulesResult:
    return rs.run_rules(office_service.process_pptx(_pptx_bytes(*boxes, **kwargs), "material.pptx"))


# --- Cálculo de contraste ---------------------------------------------------


@pytest.mark.parametrize("color, background, expected", [
    (BLACK, WHITE, 21.0),
    (WHITE, BLACK, 21.0),
    (LIGHT_GRAY, WHITE, 1.61),
    (DARK_GRAY, WHITE, 4.54),
    (WHITE, YELLOW, 1.07),
    (WHITE, WHITE, 1.0),
])
def test_relacion_de_contraste_wcag(color, background, expected):
    assert round(contrast_ratio(color, background), 2) == expected


@pytest.mark.parametrize("size, bold, expected", [
    (18, False, True), (17.5, False, False), (14, True, True), (13, True, False), (None, True, False),
])
def test_texto_grande_segun_wcag(size, bold, expected):
    assert is_large_text(size, bold) is expected


# --- PDF --------------------------------------------------------------------


def test_pdf_gris_claro_sobre_blanco_es_barrera():
    result = _pdf((TEXT * 3, LIGHT_GRAY, 11, None))

    barriers = _contrast(result)
    assert len(barriers) == 1
    barrier = barriers[0]
    assert barrier["severity"] == "barrera_detectada"
    assert barrier["category"] == "accesibilidad_visual"
    assert barrier["source"] == "rules"
    assert "1,6:1" in barrier["issue"] and "(página 1)" in barrier["issue"]
    assert "#CCCCCC" in barrier["explanation"] and "#FFFFFF" in barrier["explanation"]
    assert "4,5:1" in barrier["recommendation"]
    assert barrier["fragment_found"] is True
    assert _visual(result).status == "barrera_detectada"
    assert _penalty(result).points == SEVERE_CONTRAST_POINTS
    assert result.metrics.low_contrast_passages == 1


def test_criterio_de_aceptacion_gris_claro_ya_no_saca_lo_mismo_que_negro():
    black = rs.compute_score(None, _pdf((TEXT * 3, BLACK, 11, None)))
    gray = rs.compute_score(None, _pdf((TEXT * 3, LIGHT_GRAY, 11, None)))
    assert gray.rules_score == black.rules_score - SEVERE_CONTRAST_POINTS
    assert gray.final_score <= rs.SEVERE_BARRIER_CAP


def test_pdf_blanco_sobre_rectangulo_amarillo():
    barriers = _contrast(_pdf((TEXT * 3, WHITE, 11, YELLOW)))
    assert len(barriers) == 1
    assert "#FFFF00" in barriers[0]["explanation"]
    assert barriers[0]["severity"] == "barrera_detectada"


def test_pdf_blanco_sobre_fondo_oscuro_no_marca():
    assert _contrast(_pdf((TEXT * 3, WHITE, 11, NAVY))) == []


def test_pdf_gris_medio_es_para_revisar_en_texto_normal_y_pasa_en_texto_grande():
    normal = _contrast(_pdf((TEXT * 3, MID_GRAY, 11, None)))
    large = _contrast(_pdf((TEXT * 2, MID_GRAY, 20, None)))

    assert [b["severity"] for b in normal] == ["para_revisar"]
    assert "texto normal" in normal[0]["explanation"]
    assert large == []


def test_pdf_texto_invisible_no_se_evalua():
    # Modo de render 3 = invisible, como la capa de texto de un PDF pasado por OCR.
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_textbox(pymupdf.Rect(50, 50, 550, 300), TEXT * 3, fontsize=11, color=(1, 1, 1), render_mode=3)
    extracted = document_service.process_pdf(doc.tobytes(), "material.pdf")
    assert extracted.legibility.contrast_measured_chars == 0
    assert _contrast(rs.run_rules(extracted)) == []


# --- Word -------------------------------------------------------------------


def test_word_gris_claro_sobre_blanco():
    barriers = _contrast(_docx((TEXT * 3, LIGHT_GRAY, {})))
    assert [b["severity"] for b in barriers] == ["barrera_detectada"]
    assert "página" not in barriers[0]["issue"]  # Word no informa páginas


def test_word_blanco_resaltado_en_amarillo():
    barriers = _contrast(_docx((TEXT * 3, WHITE, {"highlight": WD_COLOR_INDEX.YELLOW})))
    assert len(barriers) == 1
    assert "#FFFF00" in barriers[0]["explanation"]


def test_word_sombreado_de_parrafo_y_de_celda():
    shaded = _contrast(_docx((TEXT * 3, WHITE, {"shading": "FFFF00"})))
    dark_shading = _contrast(_docx((TEXT * 3, WHITE, {"shading": "000080"})))
    cell = _contrast(_docx(cell=(TEXT * 2, WHITE, "FFFF99")))

    assert len(shaded) == 1
    assert dark_shading == []
    assert len(cell) == 1


def _docx_with_color_xml(color_xml: str) -> ExtractedDocument:
    document = Document()
    run = document.add_paragraph().add_run(TEXT * 3)
    run._r.get_or_add_rPr().append(parse_xml(f'<w:color {nsdecls("w")} {color_xml}/>'))
    buffer = io.BytesIO()
    document.save(buffer)
    return office_service.process_docx(buffer.getvalue(), "material.docx")


def test_word_color_del_tema_usa_el_valor_calculado_que_guarda_word():
    # Así guarda Word un color del tema aclarado: "Texto 1, más claro 75%" con el color ya calculado.
    light = _docx_with_color_xml('w:val="BFBFBF" w:themeColor="text1" w:themeTint="40"')
    dark = _docx_with_color_xml('w:val="2F5496" w:themeColor="accent1" w:themeShade="BF"')

    assert [b["severity"] for b in _contrast(rs.run_rules(light))] == ["barrera_detectada"]
    assert dark.legibility.contrast_measured_chars > 0
    assert _contrast(rs.run_rules(dark)) == []


def test_word_color_del_tema_sin_valor_calculado_no_se_evalua():
    extracted = _docx_with_color_xml('w:val="auto" w:themeColor="background1"')
    assert extracted.legibility.contrast_measured_chars == 0


def test_word_color_automatico_no_se_evalua():
    # Sin color explícito ("Automático"), Word elige negro o blanco según el fondo.
    document = Document()
    paragraph = document.add_paragraph(TEXT * 3)
    paragraph._p.get_or_add_pPr().append(_shd("000000"))
    buffer = io.BytesIO()
    document.save(buffer)

    extracted = office_service.process_docx(buffer.getvalue(), "material.docx")
    assert extracted.legibility.contrast_measured_chars == 0
    assert _contrast(rs.run_rules(extracted)) == []


# --- PowerPoint -------------------------------------------------------------


def test_powerpoint_blanco_sobre_forma_amarilla():
    barriers = _contrast(_pptx((TEXT * 2, WHITE, 20, YELLOW)))
    assert len(barriers) == 1
    assert "(diapositiva 1)" in barriers[0]["issue"]
    assert "texto grande" in barriers[0]["explanation"]


def test_powerpoint_usa_el_fondo_de_la_diapositiva():
    assert len(_contrast(_pptx((TEXT * 2, LIGHT_GRAY, 20, None), slide_background=WHITE))) == 1
    assert _contrast(_pptx((TEXT * 2, WHITE, 20, None), slide_background=NAVY)) == []


def test_powerpoint_resuelve_el_fondo_del_tema():
    # El fondo del patrón es "Fondo 1" del tema (bgRef), que en la plantilla por defecto es blanco.
    extracted = office_service.process_pptx(_pptx_bytes((TEXT * 2, LIGHT_GRAY, 20, None)), "material.pptx")
    assert extracted.legibility.low_contrast_passages[0].background_color == "#FFFFFF"
    assert len(_contrast(rs.run_rules(extracted))) == 1


# --- PowerPoint: colores del tema ----------------------------------------------

_A = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'


def _pptx_with(text_color_xml: str | None = None, *, slide_xml=None, shape=None, size=20) -> ExtractedDocument:
    """Diapositiva en blanco con un cuadro de texto; text_color_xml va dentro de un a:solidFill."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    if slide_xml:
        # La diapositiva nueva ya trae un clrMapOvr (masterClrMapping): se reemplaza.
        new = parse_xml(slide_xml)
        old = slide._element.find(new.tag)
        if old is not None:
            old.addprevious(new)
            old.getparent().remove(old)
        else:
            slide._element.append(new)
    box = shape(slide) if shape else slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(3))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = TEXT * 2
    if size:
        run.font.size = PptPt(size)
    if text_color_xml:
        r_pr = run._r.get_or_add_rPr()
        r_pr.append(parse_xml(f"<a:solidFill {_A}>{text_color_xml}</a:solidFill>"))
    buffer = io.BytesIO()
    presentation.save(buffer)
    return office_service.process_pptx(buffer.getvalue(), "material.pptx")


def test_powerpoint_texto_sin_color_hereda_texto_1_del_tema_y_no_marca():
    extracted = _pptx_with(None)
    assert extracted.legibility.contrast_measured_chars > 0  # negro (Texto 1) sobre blanco (Fondo 1)
    assert _contrast(rs.run_rules(extracted)) == []


@pytest.mark.parametrize("color_xml, flagged", [
    pytest.param('<a:schemeClr val="bg1"/>', True, id="fondo-1-sobre-fondo-1"),
    pytest.param('<a:schemeClr val="accent1"><a:lumMod val="20000"/><a:lumOff val="80000"/></a:schemeClr>',
                 True, id="enfasis-1-mas-claro-80"),
    pytest.param('<a:schemeClr val="tx2"/>', False, id="texto-2-azul-oscuro"),
    pytest.param('<a:schemeClr val="accent1"><a:lumMod val="50000"/></a:schemeClr>', False, id="enfasis-1-mas-oscuro-50"),
])
def test_powerpoint_colores_del_tema_con_aclarado_y_oscurecido(color_xml, flagged):
    barriers = _contrast(rs.run_rules(_pptx_with(color_xml)))
    assert bool(barriers) is flagged


def test_powerpoint_redefinicion_de_colores_de_la_diapositiva():
    # La diapositiva invierte el clrMap (tema oscuro): Fondo 1 = negro y Texto 1 = blanco.
    dark_map = (
        f'<p:clrMapOvr xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" {_A}>'
        '<a:overrideClrMapping bg1="dk1" tx1="lt1" bg2="dk2" tx2="lt2" accent1="accent1" accent2="accent2" '
        'accent3="accent3" accent4="accent4" accent5="accent5" accent6="accent6" hlink="hlink" folHlink="folHlink"/>'
        "</p:clrMapOvr>"
    )
    inherited = _pptx_with(None, slide_xml=dark_map)
    dark_text = _pptx_with('<a:srgbClr val="333333"/>', slide_xml=dark_map)

    assert inherited.legibility.contrast_measured_chars > 0
    assert _contrast(rs.run_rules(inherited)) == []  # blanco sobre negro
    assert dark_text.legibility.low_contrast_passages[0].background_color == "#000000"


def test_powerpoint_titulo_del_marcador_hereda_el_estilo_del_patron():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])  # título y contenido
    slide.shapes.title.text = "Ecosistemas acuáticos y sus seres vivos"
    slide.placeholders[1].text = TEXT * 2
    buffer = io.BytesIO()
    presentation.save(buffer)
    extracted = office_service.process_pptx(buffer.getvalue(), "material.pptx")
    assert extracted.legibility.contrast_measured_chars > 0
    assert _contrast(rs.run_rules(extracted)) == []


def test_powerpoint_forma_con_relleno_degradado_del_estilo_no_se_evalua():
    # Un rectángulo insertado toma el relleno del estilo (fillRef 3 = degradado en el tema por defecto).
    from pptx.enum.shapes import MSO_SHAPE

    extracted = _pptx_with(
        '<a:srgbClr val="FFFFFF"/>',
        shape=lambda s: s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.5), Inches(9), Inches(3)),
    )
    assert extracted.legibility.contrast_measured_chars == 0


def test_powerpoint_celda_de_tabla_sin_relleno_propio_no_se_evalua():
    # El relleno sale del estilo de tabla, que no se resuelve.
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table = slide.shapes.add_table(1, 1, Inches(0.5), Inches(0.5), Inches(9), Inches(2)).table
    run = table.cell(0, 0).text_frame.paragraphs[0].add_run()
    run.text = TEXT * 2
    run.font.color.rgb = PptRGB(*LIGHT_GRAY)
    buffer = io.BytesIO()
    presentation.save(buffer)
    extracted = office_service.process_pptx(buffer.getvalue(), "material.pptx")
    assert extracted.legibility.contrast_measured_chars == 0


def test_powerpoint_sin_tamano_solo_marca_lo_que_falla_incluso_como_texto_grande():
    # Sin tamaño conocido no se sabe si es texto grande: 3,5:1 no se marca, 1,6:1 sí.
    assert _contrast(_pptx((TEXT * 2, MID_GRAY, None, WHITE))) == []
    assert len(_contrast(_pptx((TEXT * 2, LIGHT_GRAY, None, WHITE)))) == 1


# --- Agregado y casos que no deben marcar nada --------------------------------


def test_se_listan_hasta_tres_pasajes_el_peor_primero_y_la_penalizacion_tiene_tope():
    result = _docx(
        (TEXT, MID_GRAY, {}), (TEXT, LIGHT_GRAY, {}), (TEXT, MID_GRAY, {}),
        (TEXT, MID_GRAY, {}), (TEXT, WHITE, {"highlight": WD_COLOR_INDEX.YELLOW}),
    )
    barriers = _contrast(result)
    assert len(barriers) == 3
    assert "1,1:1" in barriers[0]["issue"]  # blanco sobre amarillo, el peor
    assert _penalty(result).count == 5
    assert _penalty(result).points == MAX_CONTRAST_PENALTY


@pytest.mark.parametrize("result", [
    pytest.param(lambda: _pdf((TEXT * 3, BLACK, 11, None)), id="pdf-negro"),
    pytest.param(lambda: _pdf((TEXT * 3, DARK_GRAY, 12, None)), id="pdf-gris-oscuro-4.5"),
    pytest.param(lambda: _docx((TEXT * 3, BLACK, {})), id="docx-negro"),
    pytest.param(lambda: _docx((TEXT * 3, None, {})), id="docx-automatico"),
    pytest.param(lambda: _pptx((TEXT * 2, BLACK, 18, WHITE)), id="pptx-negro"),
])
def test_negro_sobre_blanco_no_genera_falsos_positivos(result):
    rules = result()
    assert _contrast(rules) == []
    assert _penalty(rules) is None


def test_texto_corto_con_bajo_contraste_no_se_marca():
    # Número de página o fuente en gris claro: menos de 5 palabras.
    result = _pdf((TEXT * 3, BLACK, 11, None), ("Página 1 de 3", LIGHT_GRAY, 9, None))
    assert _contrast(result) == []


def test_texto_pegado_no_se_evalua_por_contraste():
    text = TEXT * 3
    doc = ExtractedDocument(
        source_type="text", raw_text=text, text_for_analysis=text,
        word_count=rs.count_words(text), character_count=len(text),
        legibility=LegibilityStats(
            low_contrast_passages=[LowContrastPassage(
                text=text, text_color="#CCCCCC", background_color="#FFFFFF", contrast_ratio=1.61,
                required_ratio=4.5, is_large_text=False, word_count=30,
            )],
            low_contrast_passages_total=1,
        ),
    )
    result = rs.run_rules(doc)
    assert _contrast(result) == []
    assert result.metrics.low_contrast_passages == 0


def test_letra_chica_y_contraste_se_suman():
    result = _pdf((TEXT * 6, LIGHT_GRAY, 6, None))
    rules = {p.rule for p in result.penalties}
    assert {"letra_chica_cuerpo", "contraste_insuficiente"} <= rules
    comment = _visual(result).comment
    assert "6 pt" in comment and "contraste" in comment


def test_si_la_lectura_tipografica_falla_el_analisis_sigue(monkeypatch, caplog):
    from app.services import legibility_extractor

    def broken(*args, **kwargs):
        raise RuntimeError("XML inesperado")

    monkeypatch.setattr(legibility_extractor, "_PptxTheme", broken)
    extracted = office_service.process_pptx(_pptx_bytes((TEXT * 2, LIGHT_GRAY, 20, WHITE)), "material.pptx")

    assert extracted.raw_text  # el texto se extrajo igual
    assert extracted.legibility.contrast_measured_chars == 0
    assert "No se pudieron leer los datos de legibilidad" in caplog.text
