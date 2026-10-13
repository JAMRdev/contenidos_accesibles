"""
Extracción de Word (.docx) y PowerPoint (.pptx) al mismo ExtractedDocument que un PDF (Issue #47).
Los archivos se generan en el test con python-docx / python-pptx; las pruebas HTTP usan ?mock=true.
"""

import io
import zipfile

import pymupdf
import pytest
from docx import Document
from docx.shared import Inches, Pt
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches as PptxInches

from app.main import app
from app.services import office_service
from app.services.document_service import document_service
from app.services.rules_service import run_rules

client = TestClient(app)

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

PARAGRAPH = (
    "Los ecosistemas acuáticos albergan una gran diversidad de seres vivos que dependen del agua "
    "para alimentarse, reproducirse y protegerse de los depredadores."
)


def _png(width: int = 300, height: int = 200) -> bytes:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, width, height), False)
    pix.set_rect(pix.irect, (40, 120, 200))
    return pix.tobytes("png")


def _docx(build) -> bytes:
    document = Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pptx(build) -> bytes:
    presentation = Presentation()
    build(presentation)
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _guia_docx(document) -> None:
    document.add_heading("Guía de Ciencias Naturales", level=1)
    document.add_paragraph(PARAGRAPH)
    document.add_heading("1. Los ecosistemas acuáticos", level=2)
    document.add_paragraph("Observá la imagen y respondé.")
    document.add_picture(io.BytesIO(_png()), width=Inches(3))
    document.add_paragraph("Leé el texto de la página 8.", style="List Bullet")
    document.add_paragraph("Respondé las preguntas.", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Animal", "Hábitat"
    table.cell(1, 0).text, table.cell(1, 1).text = "Pez", "Río"
    document.add_picture(io.BytesIO(_png()), width=Inches(2))
    document.inline_shapes[1]._inline.docPr.set("descr", "Mapa de los ríos de la provincia")
    # Ícono diminuto: cuenta como imagen pero no se envía a la IA.
    document.add_picture(io.BytesIO(_png(80, 80)), width=Pt(20))


def _clase_pptx(presentation) -> None:
    portada = presentation.slides.add_slide(presentation.slide_layouts[0])
    portada.shapes.title.text = "El ciclo del agua"
    portada.placeholders[1].text = "Ciencias Naturales, 5.º grado"

    contenido = presentation.slides.add_slide(presentation.slide_layouts[1])
    contenido.shapes.title.text = "Estados del agua"
    body = contenido.placeholders[1].text_frame
    body.text = "Sólido: hielo"
    body.add_paragraph().text = "Líquido: agua"
    body.add_paragraph().text = "Gaseoso: vapor"

    imagenes = presentation.slides.add_slide(presentation.slide_layouts[5])
    imagenes.shapes.title.text = "Observá el esquema"
    imagenes.shapes.add_picture(io.BytesIO(_png()), PptxInches(1), PptxInches(5), width=PptxInches(3))
    con_alt = imagenes.shapes.add_picture(io.BytesIO(_png()), PptxInches(5), PptxInches(5), width=PptxInches(3))
    con_alt._element.nvPicPr.cNvPr.set("descr", "Nube sobre una montaña")

    oculta = presentation.slides.add_slide(presentation.slide_layouts[1])
    oculta.shapes.title.text = "Respuestas para el docente"
    oculta._element.set("show", "0")


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------

def test_docx_extrae_titulos_listas_tablas_e_imagenes():
    doc = document_service.process_file(_docx(_guia_docx), "guia.docx")

    assert doc.source_type == "docx"
    assert doc.filename == "guia.docx"
    assert [(h.level, h.text) for h in doc.structure.headings] == [
        (1, "Guía de Ciencias Naturales"),
        (2, "1. Los ecosistemas acuáticos"),
    ]
    assert doc.structure.hierarchy_levels == 2
    assert doc.structure.list_items_count == 2
    assert "• Leé el texto de la página 8.\n• Respondé las preguntas." in doc.raw_text
    assert "Animal | Hábitat\nPez | Río" in doc.raw_text
    assert PARAGRAPH in doc.raw_text

    assert doc.image_count == 3
    assert doc.images_with_alt_text == 1
    # Solo la imagen relevante sin texto alternativo va a la IA.
    assert len(doc.extracted_images) == 1
    image = doc.extracted_images[0]
    assert image.image_bytes.startswith(b"\x89PNG")
    assert image.location_description == "Imagen 1, sección «1. Los ecosistemas acuáticos»"
    assert "Observá la imagen y respondé." in image.context_text
    assert any("1 ya tiene(n) texto alternativo" in w for w in doc.warnings)


def test_docx_imagen_decorativa_no_pide_texto_alternativo():
    def build(document):
        document.add_paragraph(PARAGRAPH)
        document.add_picture(io.BytesIO(_png()), width=Inches(3))
        doc_pr = document.inline_shapes[0]._inline.docPr
        ext_lst = doc_pr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}extLst", {})
        ext = ext_lst.makeelement(
            "{http://schemas.openxmlformats.org/drawingml/2006/main}ext",
            {"uri": "{C183D7F6-B498-43B3-948B-1728B52AA6E4}"},
        )
        ext.append(ext.makeelement("{http://schemas.microsoft.com/office/drawing/2017/decorative}decorative", {"val": "1"}))
        ext_lst.append(ext)
        doc_pr.append(ext_lst)

    doc = document_service.process_file(_docx(build), "decorativa.docx")

    assert doc.image_count == 1
    assert doc.images_with_alt_text == 1
    assert doc.extracted_images == []


def test_docx_imagen_enorme_en_pixeles_no_se_decodifica(monkeypatch):
    monkeypatch.setattr(office_service, "MAX_IMAGE_PIXELS", 1000)

    doc = document_service.process_file(_docx(_guia_docx), "guia.docx")

    assert doc.extracted_images == []
    assert any(
        "No se pudo procesar automáticamente la imagen (Imagen 1, sección «1. Los ecosistemas acuáticos»)" in w
        for w in doc.warnings
    )


def test_docx_sin_estilos_de_titulo_no_tiene_encabezados():
    doc = document_service.process_file(_docx(lambda d: d.add_paragraph(PARAGRAPH)), "plano.docx")

    assert doc.structure.has_headings is False
    assert doc.raw_text == PARAGRAPH
    assert doc.page_count == 1


# ---------------------------------------------------------------------------
# PowerPoint
# ---------------------------------------------------------------------------

def test_pptx_extrae_titulos_vinetas_e_imagenes_y_omite_ocultas():
    doc = document_service.process_file(_pptx(_clase_pptx), "clase.pptx")

    assert doc.source_type == "pptx"
    assert doc.page_count == 3  # la diapositiva oculta no cuenta
    assert [(h.level, h.text, h.page_number) for h in doc.structure.headings] == [
        (1, "El ciclo del agua", 1),
        (2, "Estados del agua", 2),
        (2, "Observá el esquema", 3),
    ]
    assert doc.structure.list_items_count == 3
    assert "• Sólido: hielo\n• Líquido: agua\n• Gaseoso: vapor" in doc.raw_text
    assert "Ciencias Naturales, 5.º grado" in doc.raw_text
    assert "Respuestas para el docente" not in doc.raw_text

    # python-pptx guarda "image.png" como descripción: un nombre de archivo no cuenta como texto alternativo.
    assert doc.image_count == 2
    assert doc.images_with_alt_text == 1
    assert len(doc.extracted_images) == 1
    image = doc.extracted_images[0]
    assert image.page_number == 3
    assert image.location_description == "Diapositiva 3, sector inferior"
    assert "Observá el esquema" in image.context_text


def test_pptx_sin_diapositivas_visibles_es_error():
    def build(presentation):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide._element.set("show", "0")

    with pytest.raises(ValueError, match="no contiene diapositivas visibles"):
        document_service.process_file(_pptx(build), "vacia.pptx")


# ---------------------------------------------------------------------------
# Detección de formato y archivos inválidos
# ---------------------------------------------------------------------------

def test_extension_que_no_coincide_con_el_contenido_es_error():
    with pytest.raises(ValueError, match=r"\(parece un archivo \.pptx\)"):
        document_service.process_file(_pptx(_clase_pptx), "clase.docx")


def test_archivo_protegido_o_formato_antiguo_es_error():
    with pytest.raises(ValueError, match="protegido con contraseña o es un formato antiguo"):
        document_service.process_file(office_service.OLE_MAGIC + b"\x00" * 512, "viejo.docx")


def test_documento_con_macros_es_error():
    source = _docx(lambda d: d.add_paragraph(PARAGRAPH))
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(source)) as src, zipfile.ZipFile(buffer, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
                    b"application/vnd.ms-word.document.macroEnabled.main+xml",
                )
            dst.writestr(item, data)

    with pytest.raises(ValueError, match=r"macros \(\.docm\)"):
        document_service.process_file(buffer.getvalue(), "macros.docx")


def test_zip_que_no_es_office_es_error():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as zf:
        zf.writestr("hola.txt", "no soy un documento")

    with pytest.raises(ValueError, match="no corresponde a un documento de Word o PowerPoint"):
        document_service.process_file(buffer.getvalue(), "falso.docx")


def test_zip_bomb_se_rechaza_antes_de_abrirlo(monkeypatch):
    monkeypatch.setattr(office_service, "MAX_UNCOMPRESSED_BYTES", 1000)

    with pytest.raises(ValueError, match="demasiado grande una vez descomprimido"):
        document_service.process_file(_docx(lambda d: d.add_paragraph(PARAGRAPH)), "grande.docx")


# ---------------------------------------------------------------------------
# Reglas: estructura e imágenes con texto alternativo
# ---------------------------------------------------------------------------

def test_reglas_reconocen_la_estructura_de_word():
    doc = document_service.process_file(_docx(_guia_docx), "guia.docx")
    rules = run_rules(doc)

    estructura = next(d for d in rules.dimensions if d.category == "estructura")
    assert estructura.status == "correcta"
    assert "2 encabezados" in estructura.comment


def test_reglas_docx_largo_sin_titulos_recomienda_estilos_de_word():
    doc = document_service.process_file(
        _docx(lambda d: [d.add_paragraph(PARAGRAPH) for _ in range(8)]), "largo.docx"
    )
    rules = run_rules(doc)

    barrier = next(b for b in (f.barrier for f in rules.findings) if b["issue"] == "Documento extenso sin títulos ni estructura definida")
    assert "estilos de título" in barrier["recommendation"]


def test_reglas_todas_las_imagenes_con_texto_alternativo_no_penalizan():
    def build(document):
        document.add_paragraph(PARAGRAPH)
        document.add_picture(io.BytesIO(_png()), width=Inches(3))
        document.inline_shapes[0]._inline.docPr.set("descr", "Esquema del ciclo del agua")

    rules = run_rules(document_service.process_file(_docx(build), "con_alt.docx"))

    visual = next(d for d in rules.dimensions if d.category == "accesibilidad_visual")
    assert visual.status == "correcta"
    assert not any(p.rule == "imagenes_sin_verificar" for p in rules.penalties)


def test_reglas_solo_cuentan_las_imagenes_sin_texto_alternativo():
    rules = run_rules(document_service.process_file(_docx(_guia_docx), "guia.docx"))

    barrier = next(b for b in (f.barrier for f in rules.findings) if b["issue"] == "Imágenes sin descripción verificable")
    assert barrier["original_text"] == "[2 imagen(es) sin texto alternativo en el documento]"
    penalty = next(p for p in rules.penalties if p.rule == "imagenes_sin_verificar")
    assert penalty.count == 2


# ---------------------------------------------------------------------------
# Endpoints (siempre ?mock=true)
# ---------------------------------------------------------------------------

def test_documents_upload_acepta_docx():
    response = client.post(
        "/api/v1/documents/upload",
        files={"file": ("guia.docx", _docx(_guia_docx), DOCX_MIME)},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["source_type"] == "docx"
    assert body["images_with_alt_text"] == 1


def test_analysis_file_mock_acepta_pptx_con_propuestas_de_alt_text():
    response = client.post(
        "/api/v1/analysis/file?mock=true",
        files={"file": ("clase.pptx", _pptx(_clase_pptx), PPTX_MIME)},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["metadata"]["source_type"] == "pptx"
    assert body["metadata"]["page_count"] == 3
    assert [p["location_description"] for p in body["alt_text_proposals"]] == ["Diapositiva 3, sector inferior"]


def test_analysis_file_mock_sigue_aceptando_pdf():
    pdf = pymupdf.open()
    pdf.new_page().insert_text((50, 80), "Actividad 1: leé el texto y respondé las preguntas.")
    response = client.post(
        "/api/v1/analysis/file?mock=true",
        files={"file": ("actividad.pdf", pdf.tobytes(), "application/pdf")},
    )

    assert response.status_code == 200
    assert response.json()["metadata"]["source_type"] == "pdf"


def test_analysis_file_docx_con_datos_personales_responde_422():
    def build(document):
        document.add_paragraph("Alumno: Juan Pérez")
        document.add_paragraph("DNI: 30.123.456")
        document.add_paragraph(PARAGRAPH)

    response = client.post(
        "/api/v1/analysis/file?mock=true",
        files={"file": ("planilla.docx", _docx(build), DOCX_MIME)},
    )

    assert response.status_code == 422
    assert response.json()["detail"]["code"] == "sensitive_data_detected"
    assert "30.123.456" not in response.text


@pytest.mark.parametrize(
    ("filename", "content", "detail"),
    [
        ("notas.txt", b"hola", "Formato no admitido"),
        ("viejo.doc", b"hola", "Formato no admitido"),
        ("guia.docx", b"%PDF-1.7 no soy word", "no corresponde a un documento de Word o PowerPoint"),
        ("guia.pdf", b"PK\x03\x04 no soy pdf", "firma '%PDF-' ausente"),
        ("vacio.docx", b"", "vacío"),
    ],
)
def test_analysis_file_rechaza_archivos_invalidos(filename, content, detail):
    response = client.post(
        "/api/v1/analysis/file?mock=true",
        files={"file": (filename, content, "application/octet-stream")},
    )

    assert response.status_code == 400
    assert detail in response.json()["detail"]


def test_analysis_file_rechaza_content_type_invalido():
    response = client.post(
        "/api/v1/analysis/file?mock=true",
        files={"file": ("guia.docx", _docx(_guia_docx), "image/png")},
    )

    assert response.status_code == 400
    assert "Content-Type" in response.json()["detail"]
