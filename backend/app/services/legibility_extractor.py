"""
Extracción de datos tipográficos para las reglas de legibilidad visual (issue #79).

Hace su propia pasada sobre el archivo ya abierto para no modificar la extracción de texto
ni la detección de títulos de document_service.py y office_service.py. Todo es local y en
memoria: no consume tokens.

Por cada formato se mide:
- body_font_size: el tamaño de letra predominante, ponderado por caracteres.
- small_text_passages: párrafos (o bloques del PDF) con muchas palabras en letra chica.
- low_contrast_passages: párrafos cuyo texto no alcanza el contraste de WCAG 1.4.3 con su fondo.
- typography_issues: cursiva, mayúsculas, subrayado, fuentes decorativas, justificado e interlineado
  (las reglas de agrupación están en legibility_typography.py).

Criterio general: un dato que no se conoce (tamaño o color heredado del patrón de diapositivas,
color "automático" o de tema en Word, fondo con degradado o imagen) no se inventa. Ese
fragmento no se evalúa: es preferible no marcar nada que marcar algo falso.
"""

import bisect
import colorsys
import functools
import logging
from dataclasses import dataclass

from docx.enum.dml import MSO_COLOR_TYPE as DOCX_COLOR_TYPE
from docx.enum.text import WD_UNDERLINE
from docx.opc.constants import RELATIONSHIP_TYPE as DOCX_RT
from docx.oxml.ns import qn as w_qn
from docx.table import Table
from lxml import etree
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn as p_qn
from pptx.shapes.group import GroupShape

from app.schemas.document import LegibilityStats, LowContrastPassage, SmallTextPassage
from app.services.legibility_typography import ParagraphFormat, TypographyCollector

# Un pasaje cuenta como "letra chica" si tiene al menos MIN_SMALL_PASSAGE_WORDS palabras en menos
# de SMALL_PASSAGE_PT. Así no se marcan números de página, pies de foto de dos palabras ni marcas
# de agua cortas.
SMALL_PASSAGE_PT = 8.0
MIN_SMALL_PASSAGE_WORDS = 20
MAX_STORED_PASSAGES = 20  # el total se informa aparte en *_passages_total
PASSAGE_EXCERPT_WORDS = 40

# Contraste (WCAG 1.4.3, nivel AA): 4.5:1 para texto normal y 3:1 para texto grande.
NORMAL_TEXT_MIN_RATIO = 4.5
LARGE_TEXT_MIN_RATIO = 3.0
LARGE_TEXT_PT = 18.0
LARGE_BOLD_TEXT_PT = 14.0
# Pasajes de menos palabras ("Página 1 de 3" en gris, una fuente corta) no se marcan por contraste.
MIN_LOW_CONTRAST_WORDS = 5

WHITE = (255, 255, 255)
RGB = tuple[int, int, int]

logger = logging.getLogger(__name__)


def _never_fails(extract):
    """
    Si la lectura tipográfica falla con un archivo inesperado, el análisis sigue sin las reglas
    de legibilidad (LegibilityStats vacío) en lugar de rechazar el material del docente.
    """
    @functools.wraps(extract)
    def wrapper(document) -> LegibilityStats:
        try:
            return extract(document)
        except Exception:  # noqa: BLE001 - cualquier error del archivo deja las reglas sin datos
            logger.warning("No se pudieron leer los datos de legibilidad (%s)", extract.__name__, exc_info=True)
            return LegibilityStats()
    return wrapper


@dataclass
class Fragment:
    """Trozo de texto con formato uniforme. None en un dato significa 'no se conoce'."""

    text: str
    size: float | None = None
    color: RGB | None = None
    background: RGB | None = None
    bold: bool = False
    italic: bool = False
    underline: bool = False
    caps: bool = False  # formato "Todo mayúsculas": se ve en mayúsculas aunque el texto no lo esté
    font: str | None = None
    link: bool = False
    visible: bool = True  # False: texto invisible del PDF (por ejemplo, la capa de OCR)


# ---------------------------------------------------------------------------
# Contraste (WCAG 2.x)
# ---------------------------------------------------------------------------

def _channel(value: int) -> float:
    c = value / 255
    return c / 12.92 if c <= 0.04045 else ((c + 0.055) / 1.055) ** 2.4


def relative_luminance(rgb: RGB) -> float:
    r, g, b = (_channel(v) for v in rgb)
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def contrast_ratio(a: RGB, b: RGB) -> float:
    """Relación de contraste entre dos colores, de 1 (iguales) a 21 (negro sobre blanco)."""
    la, lb = sorted((relative_luminance(a), relative_luminance(b)), reverse=True)
    return (la + 0.05) / (lb + 0.05)


def is_large_text(size: float | None, bold: bool) -> bool:
    return size is not None and (size >= LARGE_TEXT_PT or (bold and size >= LARGE_BOLD_TEXT_PT))


def hex_color(rgb: RGB) -> str:
    return "#{:02X}{:02X}{:02X}".format(*rgb)


def parse_hex(value: str | None) -> RGB | None:
    """'CCCCCC' o '#CCCCCC' a (204, 204, 204). 'auto' o un valor inválido devuelven None."""
    value = (value or "").lstrip("#")
    if len(value) != 6:
        return None
    try:
        return tuple(int(value[i:i + 2], 16) for i in (0, 2, 4))
    except ValueError:
        return None


def _excerpt(texts: list[str]) -> str:
    words = " ".join(texts).split()
    return " ".join(words[:PASSAGE_EXCERPT_WORDS]) + ("…" if len(words) > PASSAGE_EXCERPT_WORDS else "")


class _Collector:
    """Acumula tamaños y contraste por caracteres, y arma las listas de pasajes."""

    def __init__(self) -> None:
        self.char_counts: dict[float, int] = {}
        self.passages: list[SmallTextPassage] = []
        self.passages_total = 0
        self.contrast_chars = 0
        self.low_contrast: list[LowContrastPassage] = []
        self.low_contrast_total = 0
        self.typography = TypographyCollector()

    def add_paragraph(
        self, fragments: list[Fragment], page_number: int, fmt: ParagraphFormat | None = None,
        has_pages: bool = True,
    ) -> None:
        """fragments: los trozos de un mismo párrafo (Word/PowerPoint) o bloque (PDF)."""
        page_number = max(page_number, 1)
        self._small_text(fragments, page_number)
        self._contrast(fragments, page_number)
        self.typography.add_paragraph(fragments, page_number if has_pages else None, fmt)

    def _small_text(self, fragments: list[Fragment], page_number: int) -> None:
        small_texts: list[str] = []
        small_sizes: dict[float, int] = {}
        small_words = 0
        for fragment in fragments:
            stripped = fragment.text.strip()
            if not stripped or fragment.size is None:
                continue
            size = round(fragment.size, 1)
            self.char_counts[size] = self.char_counts.get(size, 0) + len(stripped)
            if size < SMALL_PASSAGE_PT:
                small_texts.append(stripped)
                small_sizes[size] = small_sizes.get(size, 0) + len(stripped)
                small_words += len(stripped.split())

        if small_words < MIN_SMALL_PASSAGE_WORDS:
            return
        self.passages_total += 1
        if len(self.passages) < MAX_STORED_PASSAGES:
            self.passages.append(SmallTextPassage(
                text=_excerpt(small_texts),
                page_number=page_number,
                font_size=max(small_sizes, key=small_sizes.get),
                word_count=small_words,
            ))

    def _contrast(self, fragments: list[Fragment], page_number: int) -> None:
        # (color, fondo, grande) -> [caracteres, palabras, textos, ratio]
        failing: dict[tuple[RGB, RGB, bool], list] = {}
        for fragment in fragments:
            stripped = fragment.text.strip()
            if not stripped or fragment.color is None or fragment.background is None:
                continue
            self.contrast_chars += len(stripped)
            ratio = contrast_ratio(fragment.color, fragment.background)
            large = is_large_text(fragment.size, fragment.bold)
            # Sin tamaño conocido no se sabe si el texto es grande: solo se marca lo que falla
            # incluso con el umbral de texto grande.
            required = LARGE_TEXT_MIN_RATIO if large or fragment.size is None else NORMAL_TEXT_MIN_RATIO
            if ratio >= required:
                continue
            key = (fragment.color, fragment.background, large)
            entry = failing.setdefault(key, [0, 0, [], ratio])
            entry[0] += len(stripped)
            entry[1] += len(stripped.split())
            entry[2].append(stripped)

        words = sum(e[1] for e in failing.values())
        if words < MIN_LOW_CONTRAST_WORDS:
            return
        self.low_contrast_total += 1
        if len(self.low_contrast) >= MAX_STORED_PASSAGES:
            return
        # Se informa la combinación de colores que más texto ocupa dentro del pasaje.
        (color, background, large), (_, _, _, ratio) = max(failing.items(), key=lambda item: item[1][0])
        self.low_contrast.append(LowContrastPassage(
            text=_excerpt([t for e in failing.values() for t in e[2]]),
            page_number=page_number,
            text_color=hex_color(color),
            background_color=hex_color(background),
            contrast_ratio=round(ratio, 2),
            required_ratio=LARGE_TEXT_MIN_RATIO if large else NORMAL_TEXT_MIN_RATIO,
            is_large_text=large,
            word_count=words,
        ))

    def result(self) -> LegibilityStats:
        body_font, body_font_decorative, typography_issues = self.typography.result()
        return LegibilityStats(
            body_font_size=max(self.char_counts, key=self.char_counts.get) if self.char_counts else None,
            measured_chars=sum(self.char_counts.values()),
            small_text_passages=self.passages,
            small_text_passages_total=self.passages_total,
            contrast_measured_chars=self.contrast_chars,
            low_contrast_passages=self.low_contrast,
            low_contrast_passages_total=self.low_contrast_total,
            body_font_name=body_font,
            body_font_decorative=body_font_decorative,
            typography_issues=typography_issues,
        )


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

MIN_BACKGROUND_FILL_OPACITY = 0.5


def _int_to_rgb(value: int) -> RGB:
    return ((value >> 16) & 0xFF, (value >> 8) & 0xFF, value & 0xFF)


def _float_to_rgb(fill) -> RGB | None:
    """Color de relleno de PyMuPDF (tupla de 0 a 1, en RGB, gris o CMYK) a RGB de 0 a 255."""
    if not fill:
        return None
    if len(fill) == 1:
        fill = (fill[0],) * 3
    elif len(fill) == 4:  # CMYK
        c, m, y, k = fill
        fill = ((1 - c) * (1 - k), (1 - m) * (1 - k), (1 - y) * (1 - k))
    elif len(fill) != 3:
        return None
    return tuple(round(max(0.0, min(1.0, v)) * 255) for v in fill)


MAX_UNDERLINE_THICKNESS = 2.0  # un rectángulo más alto que esto es un relleno, no un subrayado


def _pdf_drawings(page) -> tuple[list[tuple], list[tuple]]:
    """
    Una sola pasada por los dibujos de la página:
    - fills: rectángulos rellenos opacos, en orden de dibujo: (rect, color).
    - segments: líneas horizontales finas (posibles subrayados), ordenadas por y: (y, x0, x1).
    """
    fills, segments = [], []
    for drawing in page.get_drawings():
        color = _float_to_rgb(drawing.get("fill"))
        opacity = drawing.get("fill_opacity")
        if color is not None and (opacity is None or opacity >= MIN_BACKGROUND_FILL_OPACITY):
            fills.append((drawing["rect"], color))
        for item in drawing.get("items", []):
            if item[0] == "l":
                p1, p2 = item[1], item[2]
                if abs(p1.y - p2.y) < 0.8 and abs(p1.x - p2.x) > 2:
                    segments.append(((p1.y + p2.y) / 2, min(p1.x, p2.x), max(p1.x, p2.x)))
            elif item[0] == "re":
                rect = item[1]
                if 0 < rect.height <= MAX_UNDERLINE_THICKNESS and rect.width > 2:
                    segments.append(((rect.y0 + rect.y1) / 2, rect.x0, rect.x1))
    segments.sort()
    return fills, segments


def _pdf_underlined(span: dict, line_bbox, segments: list[tuple]) -> bool:
    """
    True si hay una línea fina justo debajo de la base del texto que cubre al menos el 70% del span
    y no es más ancha que el renglón (así no se confunde con el borde de una celda de tabla).
    """
    x0, _, x1, _ = span["bbox"]
    width = x1 - x0
    if width <= 0 or not segments:
        return False
    baseline = span.get("origin", (0, span["bbox"][3]))[1]
    size = span.get("size") or 10
    low, high = baseline - 1.0, baseline + 0.3 * size
    start = bisect.bisect_left(segments, (low, float("-inf"), float("-inf")))
    for y, sx0, sx1 in segments[start:]:
        if y > high:
            break
        overlap = min(x1, sx1) - max(x0, sx0)
        if overlap >= 0.7 * width and sx0 >= line_bbox[0] - 4 and sx1 <= line_bbox[2] + 4:
            return True
    return False


def _pdf_in_link(bbox, links: list) -> bool:
    x = (bbox[0] + bbox[2]) / 2
    y = (bbox[1] + bbox[3]) / 2
    return any(r.x0 <= x <= r.x1 and r.y0 <= y <= r.y1 for r in links)


def _pdf_background(bbox, fills: list[tuple]) -> RGB:
    """Fondo del texto: el último rectángulo relleno que contiene su centro; si no hay, blanco."""
    x = (bbox[0] + bbox[2]) / 2
    y = (bbox[1] + bbox[3]) / 2
    for rect, color in reversed(fills):
        if rect.x0 <= x <= rect.x1 and rect.y0 <= y <= rect.y1:
            return color
    return WHITE


@_never_fails
def extract_pdf_legibility(doc) -> LegibilityStats:
    """doc es un pymupdf.Document ya abierto. Cada bloque de texto se trata como un párrafo."""
    collector = _Collector()
    for page_idx in range(len(doc)):
        page = doc.load_page(page_idx)
        fills, segments = _pdf_drawings(page)
        links = [link["from"] for link in page.get_links() if link.get("from") is not None]
        for block in page.get_text("dict").get("blocks", []):
            if block.get("type") != 0:  # 0 = texto
                continue
            fragments = []
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    # alpha 0: texto invisible (por ejemplo, la capa de OCR de un escaneo).
                    visible = span.get("alpha", 255) > 0
                    flags = span.get("flags", 0)
                    font = (span.get("font") or None) if visible else None
                    fragments.append(Fragment(
                        text=span.get("text", ""),
                        size=span.get("size"),
                        color=_int_to_rgb(span.get("color", 0)) if visible else None,
                        background=_pdf_background(span["bbox"], fills) if visible else None,
                        bold=bool(flags & 16),
                        italic=bool(flags & 2) or any(w in (font or "").lower() for w in ("italic", "oblique")),
                        underline=visible and _pdf_underlined(span, line["bbox"], segments),
                        font=font,
                        link=_pdf_in_link(span["bbox"], links),
                        visible=visible,
                    ))
            collector.add_paragraph(fragments, page_idx + 1)
    return collector.result()


# ---------------------------------------------------------------------------
# Word (.docx)
# ---------------------------------------------------------------------------

# Colores de resaltado de Word (w:highlight).
_HIGHLIGHT_COLORS = {
    "yellow": "FFFF00", "green": "00FF00", "cyan": "00FFFF", "magenta": "FF00FF",
    "blue": "0000FF", "red": "FF0000", "darkBlue": "000080", "darkCyan": "008080",
    "darkGreen": "008000", "darkMagenta": "800080", "darkRed": "800000",
    "darkYellow": "808000", "darkGray": "808080", "lightGray": "C0C0C0",
    "black": "000000", "white": "FFFFFF",
}


def _style_chain(style):
    while style is not None:
        yield style
        style = style.base_style


def _style_sizes(style):
    """Tamaño definido en el estilo o en alguno de sus estilos base."""
    for s in _style_chain(style):
        font = getattr(s, "font", None)
        if font is not None and font.size is not None:
            return font.size.pt
    return None


def _docx_default_size(document) -> float | None:
    """Tamaño de docDefaults (w:sz está en medios puntos). Word usa 10 pt si no está definido."""
    sz = document.styles.element.find(
        f"{w_qn('w:docDefaults')}/{w_qn('w:rPrDefault')}/{w_qn('w:rPr')}/{w_qn('w:sz')}"
    )
    if sz is not None and sz.get(w_qn("w:val"), "").isdigit():
        return int(sz.get(w_qn("w:val"))) / 2
    return 10.0


def _docx_run_size(run, paragraph, default_size: float | None) -> float | None:
    if run.font.size is not None:
        return run.font.size.pt
    for style in (run.style, paragraph.style):
        size = _style_sizes(style) if style is not None else None
        if size is not None:
            return size
    return default_size


def _rgb_of(font_color) -> tuple[bool, RGB | None]:
    """
    (definido, color). Con un color del tema, Word guarda además en w:val el color ya calculado
    (con su aclarado u oscurecido), así que se usa ese. Definido sin color ("Automático", o un
    color del tema sin w:val): no se puede saber cómo se ve.
    """
    if font_color is None or font_color.type is None:
        return False, None
    if font_color.type not in (DOCX_COLOR_TYPE.RGB, DOCX_COLOR_TYPE.THEME):
        return True, None
    rgb = font_color.rgb
    return True, tuple(rgb) if rgb is not None else None


def _docx_run_color(run, paragraph) -> RGB | None:
    """Color explícito del run o heredado de sus estilos. Automático o de tema: None (no se evalúa)."""
    defined, rgb = _rgb_of(run.font.color)
    if defined:
        return rgb
    for style in (run.style, paragraph.style):
        for s in _style_chain(style):
            font = getattr(s, "font", None)
            if font is None:
                continue
            defined, rgb = _rgb_of(font.color)
            if defined:
                return rgb
    return None


def _shading_fill(properties) -> RGB | None:
    """Relleno de w:shd dentro de rPr, pPr o tcPr ('auto' o ausente: None)."""
    if properties is None:
        return None
    shd = properties.find(w_qn("w:shd"))
    return parse_hex(shd.get(w_qn("w:fill"))) if shd is not None else None


def _docx_run_background(run, paragraph_fill: RGB | None) -> RGB | None:
    r_pr = run._r.rPr
    if r_pr is not None:
        highlight = r_pr.find(w_qn("w:highlight"))
        if highlight is not None and highlight.get(w_qn("w:val")) in _HIGHLIGHT_COLORS:
            return parse_hex(_HIGHLIGHT_COLORS[highlight.get(w_qn("w:val"))])
        run_fill = _shading_fill(r_pr)
        if run_fill is not None:
            return run_fill
    return paragraph_fill


def _docx_page_background(document) -> RGB:
    """Color de página (Diseño > Color de página); blanco si no está definido."""
    background = document.element.find(w_qn("w:background"))
    color = parse_hex(background.get(w_qn("w:color"))) if background is not None else None
    return color or WHITE


# --- Word: tipografía (issue #79, parte 3) ---

_JUSTIFIED_JC = {"both", "distribute", "thaiDistribute", "lowKashida", "mediumKashida", "highKashida"}
_HYPERLINK_STYLES = ("hyperlink", "followedhyperlink", "hipervínculo", "hipervinculo")


def _theme_fonts(theme_blob: bytes | None) -> dict[str, str]:
    """Fuentes latinas del tema: {'major': 'Calibri Light', 'minor': 'Calibri'}."""
    fonts: dict[str, str] = {}
    if not theme_blob:
        return fonts
    try:
        theme = etree.fromstring(theme_blob)
    except etree.XMLSyntaxError:
        return fonts
    for kind in ("major", "minor"):
        latin = theme.find(f".//a:fontScheme/a:{kind}Font/a:latin", _A_NS)
        if latin is not None and latin.get("typeface"):
            fonts[kind] = latin.get("typeface")
    return fonts


def _docx_theme_fonts(document) -> dict[str, str]:
    try:
        return _theme_fonts(document.part.part_related_by(DOCX_RT.THEME).blob)
    except KeyError:
        return {}


def _docx_rfonts(r_pr, theme_fonts: dict[str, str]) -> str | None:
    """Fuente de un rPr: w:ascii explícita o la del tema (w:asciiTheme = majorHAnsi, minorHAnsi…)."""
    if r_pr is None:
        return None
    r_fonts = r_pr.find(w_qn("w:rFonts"))
    if r_fonts is None:
        return None
    if r_fonts.get(w_qn("w:ascii")):
        return r_fonts.get(w_qn("w:ascii"))
    theme = r_fonts.get(w_qn("w:asciiTheme")) or r_fonts.get(w_qn("w:hAnsiTheme"))
    if theme:
        return theme_fonts.get("major" if theme.startswith("major") else "minor")
    return None


def _docx_style_elements(run, paragraph):
    """Estilos que aplican al run, del más específico al más general."""
    for style in (run.style, paragraph.style):
        if style is not None:
            yield from _style_chain(style)


def _docx_font_name(run, paragraph, theme_fonts: dict[str, str], default_font: str | None) -> str | None:
    name = _docx_rfonts(run._r.rPr, theme_fonts)
    if name:
        return name
    for style in _docx_style_elements(run, paragraph):
        name = _docx_rfonts(style.element.rPr, theme_fonts)
        if name:
            return name
    return default_font


def _docx_flag(run, paragraph, attribute: str) -> bool:
    """italic, all_caps: el valor del run o, si no lo define, el del primer estilo que lo define."""
    value = getattr(run.font, attribute)
    if value is not None:
        return bool(value)
    for style in _docx_style_elements(run, paragraph):
        font = getattr(style, "font", None)
        value = getattr(font, attribute, None) if font is not None else None
        if value is not None:
            return bool(value)
    return False


def _is_underlined(value) -> bool:
    return value not in (None, False) and value != WD_UNDERLINE.NONE


def _docx_underline(run, paragraph) -> bool:
    if run.font.underline is not None:
        return _is_underlined(run.font.underline)
    for style in _docx_style_elements(run, paragraph):
        font = getattr(style, "font", None)
        if font is not None and font.underline is not None:
            return _is_underlined(font.underline)
    return False


def _docx_is_link(run) -> bool:
    """Run con el estilo de carácter de hipervínculo (los runs dentro de w:hyperlink no llegan acá)."""
    style = run.style
    if style is None:
        return False
    names = {(style.style_id or "").lower(), (style.name or "").lower()}
    return any(n in _HYPERLINK_STYLES for n in names)


def _docx_paragraph_properties(paragraph, document):
    """pPr del párrafo, de su estilo y estilos base, y el de docDefaults."""
    if paragraph._p.pPr is not None:
        yield paragraph._p.pPr
    for style in _style_chain(paragraph.style):
        if style.element.pPr is not None:
            yield style.element.pPr
    defaults = document.styles.element.find(
        f"{w_qn('w:docDefaults')}/{w_qn('w:pPrDefault')}/{w_qn('w:pPr')}"
    )
    if defaults is not None:
        yield defaults


def _docx_paragraph_format(paragraph, document, font_size: float | None) -> ParagraphFormat:
    justified = None
    spacing = None
    for p_pr in _docx_paragraph_properties(paragraph, document):
        if justified is None:
            jc = p_pr.find(w_qn("w:jc"))
            if jc is not None and jc.get(w_qn("w:val")):
                justified = jc.get(w_qn("w:val")) in _JUSTIFIED_JC
        if spacing is None:
            element = p_pr.find(w_qn("w:spacing"))
            if element is not None and element.get(w_qn("w:line"), "").lstrip("-").isdigit():
                spacing = element
        if justified is not None and spacing is not None:
            break

    line_spacing = 1.0  # sin interlineado definido, Word usa "Sencillo"
    if spacing is not None:
        line = int(spacing.get(w_qn("w:line")))
        rule = spacing.get(w_qn("w:lineRule"), "auto")
        if rule == "auto":
            line_spacing = line / 240  # en 240avos de renglón
        elif font_size:
            points = line / 20  # en veinteavos de punto
            # "Mínimo": nunca queda más apretado que sencillo.
            line_spacing = points / font_size if rule == "exact" else max(1.0, points / font_size)
        else:
            line_spacing = None
    return ParagraphFormat(justified=bool(justified), line_spacing=line_spacing)


def _predominant_size(fragments: list[Fragment]) -> float | None:
    sizes: dict[float, int] = {}
    for fragment in fragments:
        if fragment.size is not None and fragment.text.strip():
            sizes[fragment.size] = sizes.get(fragment.size, 0) + len(fragment.text.strip())
    return max(sizes, key=sizes.get) if sizes else None


def _docx_paragraphs(document):
    """(párrafo, relleno de la celda) del cuerpo y de las tablas (celdas combinadas, una sola vez)."""
    for item in document.iter_inner_content():
        if isinstance(item, Table):
            seen = set()
            for row in item.rows:
                for cell in row.cells:
                    if id(cell._tc) in seen:
                        continue
                    seen.add(id(cell._tc))
                    cell_fill = _shading_fill(cell._tc.tcPr)
                    for paragraph in cell.paragraphs:
                        yield paragraph, cell_fill
        else:
            yield item, None


@_never_fails
def extract_docx_legibility(document) -> LegibilityStats:
    """document es un docx.Document ya abierto. Word no informa páginas: todo va como página 1."""
    collector = _Collector()
    default_size = _docx_default_size(document)
    page_background = _docx_page_background(document)
    theme_fonts = _docx_theme_fonts(document)
    defaults_rpr = document.styles.element.find(
        f"{w_qn('w:docDefaults')}/{w_qn('w:rPrDefault')}/{w_qn('w:rPr')}"
    )
    default_font = _docx_rfonts(defaults_rpr, theme_fonts) or theme_fonts.get("minor")
    for paragraph, cell_fill in _docx_paragraphs(document):
        paragraph_fill = _shading_fill(paragraph._p.pPr) or cell_fill or page_background
        style_bold = any(
            getattr(s, "font", None) is not None and s.font.bold for s in _style_chain(paragraph.style)
        )
        fragments = [
            Fragment(
                text=run.text,
                size=_docx_run_size(run, paragraph, default_size),
                color=_docx_run_color(run, paragraph),
                background=_docx_run_background(run, paragraph_fill),
                bold=run.bold if run.bold is not None else style_bold,
                italic=_docx_flag(run, paragraph, "italic"),
                underline=_docx_underline(run, paragraph),
                caps=_docx_flag(run, paragraph, "all_caps"),
                font=_docx_font_name(run, paragraph, theme_fonts, default_font),
                link=_docx_is_link(run),
            )
            for run in paragraph.runs
        ]
        fmt = _docx_paragraph_format(paragraph, document, _predominant_size(fragments))
        collector.add_paragraph(fragments, 1, fmt, has_pages=False)
    return collector.result()


# ---------------------------------------------------------------------------
# PowerPoint (.pptx)
# ---------------------------------------------------------------------------
#
# PowerPoint casi nunca guarda colores RGB: usa colores del tema ("Texto 1", "Énfasis 2, más claro
# 40%") que se resuelven así: el patrón traduce el nombre (clrMap: tx1 -> dk1), el tema da el
# color base y los modificadores (lumMod, lumOff, tint, shade) lo aclaran u oscurecen. El texto que
# no define color lo hereda de su lista de estilos, del marcador de posición del diseño y del patrón,
# de los estilos de texto del patrón o del estilo por defecto de la presentación.

# Nombres de clrMap que se traducen; los demás (dk1, lt1, accent1…) son colores del tema directos.
_MAPPED_SCHEME_NAMES = (
    "bg1", "tx1", "bg2", "tx2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6",
    "hlink", "folHlink",
)
_COLOR_NAMES = {"srgbClr", "schemeClr", "sysClr", "scrgbClr", "prstClr", "hslClr"}
_A_NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
_TITLE_PLACEHOLDERS = ("title", "ctrTitle")
_BODY_PLACEHOLDERS = ("body", "subTitle", "obj", None)  # sin tipo = contenido (obj)


def _pptx_font_scale(text_frame) -> float:
    """Factor de 'Reducir el texto al desbordar' (normAutofit fontScale, en milésimas de %)."""
    autofit = text_frame._bodyPr.find(p_qn("a:normAutofit"))
    scale = autofit.get("fontScale") if autofit is not None else None
    return int(scale) / 100_000 if scale and scale.isdigit() else 1.0


def _apply_modifiers(rgb: RGB, color_el) -> RGB:
    """Aplica lumMod/lumOff (en HSL, como PowerPoint), tint (hacia blanco) y shade (hacia negro)."""
    mods = {_localname(m): int(m.get("val", "100000")) / 100_000 for m in color_el}
    r, g, b = (v / 255 for v in rgb)
    if "lumMod" in mods or "lumOff" in mods:
        h, lum, s = colorsys.rgb_to_hls(r, g, b)
        lum = max(0.0, min(1.0, lum * mods.get("lumMod", 1.0) + mods.get("lumOff", 0.0)))
        r, g, b = colorsys.hls_to_rgb(h, lum, s)
    if "tint" in mods:
        r, g, b = (c * mods["tint"] + (1 - mods["tint"]) for c in (r, g, b))
    if "shade" in mods:
        r, g, b = (c * mods["shade"] for c in (r, g, b))
    return tuple(round(max(0.0, min(1.0, c)) * 255) for c in (r, g, b))


def _localname(element) -> str:
    return etree.QName(element).localname


class _PptxTheme:
    """Colores y estilos de relleno del tema de un patrón, con su clrMap."""

    def __init__(self, master) -> None:
        self.scheme: dict[str, RGB] = {}
        self.fill_styles: list = []
        self.bg_fill_styles: list = []
        try:
            theme = etree.fromstring(master.part.part_related_by(RT.THEME).blob)
        except (KeyError, etree.XMLSyntaxError):
            theme = None
        if theme is not None:
            for slot in theme.iterfind(".//a:clrScheme/*", _A_NS):
                base = slot[0] if len(slot) else None
                if base is None:
                    continue
                value = base.get("lastClr") if _localname(base) == "sysClr" else base.get("val")
                rgb = parse_hex(value)
                if rgb is not None:
                    self.scheme[_localname(slot)] = rgb
            fmt = theme.find(".//a:fmtScheme", _A_NS)
            if fmt is not None:
                fills = fmt.find("a:fillStyleLst", _A_NS)
                bg_fills = fmt.find("a:bgFillStyleLst", _A_NS)
                self.fill_styles = list(fills) if fills is not None else []
                self.bg_fill_styles = list(bg_fills) if bg_fills is not None else []
        clr_map = master._element.find(p_qn("p:clrMap"))
        self.master_map = dict(clr_map.attrib) if clr_map is not None else {}
        try:
            self.fonts = _theme_fonts(master.part.part_related_by(RT.THEME).blob)
        except KeyError:
            self.fonts = {}

    def font(self, typeface: str | None) -> str | None:
        """'+mj-lt' y '+mn-lt' son las fuentes del tema (títulos y cuerpo); el resto, tal cual."""
        if not typeface:
            return None
        if typeface.startswith("+mj"):
            return self.fonts.get("major")
        if typeface.startswith("+mn"):
            return self.fonts.get("minor")
        return typeface

    @staticmethod
    def color_map(master_map: dict, *parts) -> dict:
        """clrMap del patrón, con la redefinición del diseño o la diapositiva si la tienen."""
        mapping = dict(master_map)
        for part in parts:
            override = part._element.find(f"{p_qn('p:clrMapOvr')}/{p_qn('a:overrideClrMapping')}")
            if override is not None:
                mapping = dict(override.attrib)
        return mapping

    def resolve(self, color_el, mapping: dict, placeholder: RGB | None = None) -> RGB | None:
        """Elemento de color de DrawingML a RGB. phClr usa el color del estilo que lo llama."""
        if color_el is None:
            return None
        kind = _localname(color_el)
        if kind == "srgbClr":
            base = parse_hex(color_el.get("val"))
        elif kind == "sysClr":
            base = parse_hex(color_el.get("lastClr"))
        elif kind == "schemeClr":
            name = color_el.get("val")
            if name == "phClr":
                base = placeholder
            else:
                if name in _MAPPED_SCHEME_NAMES:
                    name = mapping.get(name, name)
                base = self.scheme.get(name)
        else:  # scrgbClr, prstClr, hslClr: poco usados en materiales; no se evalúan
            base = None
        if base is None:
            return None
        if any(_localname(m) == "alpha" and int(m.get("val", "100000")) < 50_000 for m in color_el):
            return None  # muy transparente: el color real depende de lo que hay detrás
        return _apply_modifiers(base, color_el)

    def fill(self, properties, mapping: dict, placeholder: RGB | None = None) -> tuple[bool, RGB | None]:
        """
        (define relleno, color) de spPr, tcPr o bgPr. noFill o sin relleno: no define (se ve lo de
        atrás). Degradado, imagen o trama: definido pero desconocido.
        """
        if properties is None:
            return False, None
        solid = properties.find(p_qn("a:solidFill"))
        if solid is not None:
            return True, self.resolve(_first_color(solid), mapping, placeholder)
        if any(properties.find(p_qn(t)) is not None for t in ("a:gradFill", "a:blipFill", "a:pattFill", "a:grpFill")):
            return True, None
        return False, None

    def style_fill(self, ref_el, mapping: dict) -> tuple[bool, RGB | None]:
        """Relleno de fillRef/bgRef: idx 1-999 en fillStyleLst, 1001+ en bgFillStyleLst; 0 = sin relleno."""
        idx = int(ref_el.get("idx", "0"))
        if idx == 0:
            return False, None
        styles, position = (self.bg_fill_styles, idx - 1001) if idx >= 1001 else (self.fill_styles, idx - 1)
        if not 0 <= position < len(styles):
            return True, None
        style = styles[position]
        color = self.resolve(_first_color(ref_el), mapping)
        if _localname(style) == "solidFill":
            return True, self.resolve(_first_color(style), mapping, placeholder=color)
        if _localname(style) == "noFill":
            return False, None
        return True, None


def _first_color(element):
    """Primer elemento de color (srgbClr, schemeClr…) dentro de solidFill, fillRef, fontRef, etc."""
    if element is None:
        return None
    return next((child for child in element if _localname(child) in _COLOR_NAMES), None)


def _pptx_slide_background(slide, theme: _PptxTheme, mapping: dict) -> RGB | None:
    """Fondo de la diapositiva, del diseño o del patrón (leído del XML, sin modificarlo)."""
    for part in (slide, slide.slide_layout, slide.slide_layout.slide_master):
        bg = part._element.find(p_qn("p:cSld")).find(p_qn("p:bg"))
        if bg is None:
            continue
        bg_pr = bg.find(p_qn("p:bgPr"))
        if bg_pr is not None:
            defined, color = theme.fill(bg_pr, mapping)
            return color if defined else None
        bg_ref = bg.find(p_qn("p:bgRef"))
        if bg_ref is not None:
            defined, color = theme.style_fill(bg_ref, mapping)
            return color if defined else None
        return None
    return None


def _placeholder_chain(shape) -> list:
    """El marcador de posición de la diapositiva, el del diseño y el del patrón (los que existan)."""
    chain = [shape]
    current = shape
    for _ in range(2):
        try:
            current = current._base_placeholder
        except (AttributeError, KeyError, ValueError):
            current = None
        if current is None:
            break
        chain.append(current)
    return chain


def _placeholder_type(shape) -> str | None:
    ph = shape._element.find(f".//{p_qn('p:nvPr')}/{p_qn('p:ph')}")
    return ph.get("type") if ph is not None else None


def _is_placeholder(shape) -> bool:
    return shape._element.find(f".//{p_qn('p:nvPr')}/{p_qn('p:ph')}") is not None


def _shape_fill(shape, theme: _PptxTheme, mapping: dict) -> tuple[bool, RGB | None]:
    """Relleno propio, heredado del marcador de posición o del estilo de la forma (p:style fillRef)."""
    chain = _placeholder_chain(shape) if _is_placeholder(shape) else [shape]
    for item in chain:
        defined, color = theme.fill(item._element.spPr, mapping)
        if defined:
            return defined, color
        if item._element.spPr is not None and item._element.spPr.find(p_qn("a:noFill")) is not None:
            return False, None
    style = shape._element.find(p_qn("p:style"))
    fill_ref = style.find(p_qn("a:fillRef")) if style is not None else None
    if fill_ref is not None:
        return theme.style_fill(fill_ref, mapping)
    return False, None


def _def_rpr(lvl_ppr):
    return lvl_ppr.find(p_qn("a:defRPr")) if lvl_ppr is not None else None


class _PptxTextDefaults:
    """
    Formato por defecto del texto de una forma según la cadena de herencia de PowerPoint: lista de
    estilos del cuadro, marcadores de posición del diseño y del patrón, estilos de texto del patrón
    (título, cuerpo u otros) y, para las formas que no son marcadores, el estilo de la forma
    (fontRef) y el estilo por defecto de la presentación.
    """

    def __init__(self, presentation, theme: _PptxTheme, master, mapping: dict) -> None:
        self.theme = theme
        self.mapping = mapping
        self.default_style = presentation._element.find(p_qn("p:defaultTextStyle"))
        tx_styles = master._element.find(p_qn("p:txStyles"))
        self.title_style = tx_styles.find(p_qn("p:titleStyle")) if tx_styles is not None else None
        self.body_style = tx_styles.find(p_qn("p:bodyStyle")) if tx_styles is not None else None
        self.other_style = tx_styles.find(p_qn("p:otherStyle")) if tx_styles is not None else None

    def _list_styles(self, shape, text_frame) -> list:
        sources = [text_frame._txBody.find(p_qn("a:lstStyle"))]
        if _is_placeholder(shape):
            for base in _placeholder_chain(shape)[1:]:
                tx_body = base._element.find(p_qn("p:txBody"))
                sources.append(tx_body.find(p_qn("a:lstStyle")) if tx_body is not None else None)
            ph_type = _placeholder_type(shape)
            sources.append(
                self.title_style if ph_type in _TITLE_PLACEHOLDERS
                else self.body_style if ph_type in _BODY_PLACEHOLDERS
                else self.other_style
            )
        return sources

    def levels(self, shape, text_frame, level: int) -> tuple[list, list]:
        """(niveles principales, niveles del estilo por defecto): elementos a:lvlNpPr que existen."""
        tag = p_qn(f"a:lvl{level + 1}pPr")
        main = [s.find(tag) for s in self._list_styles(shape, text_frame) if s is not None]
        default = []
        if not _is_placeholder(shape) and self.default_style is not None:
            default = [self.default_style.find(tag)]
        return [lvl for lvl in main if lvl is not None], [lvl for lvl in default if lvl is not None]

    def _font_ref(self, shape):
        if _is_placeholder(shape):
            return None
        style = shape._element.find(p_qn("p:style"))
        return style.find(p_qn("a:fontRef")) if style is not None else None

    def color(self, shape, text_frame, level: int) -> RGB | None:
        main, default = self.levels(shape, text_frame, level)
        for lvl in main:
            solid = _def_rpr(lvl).find(p_qn("a:solidFill")) if _def_rpr(lvl) is not None else None
            if solid is not None:
                return self.theme.resolve(_first_color(solid), self.mapping)
        font_ref = self._font_ref(shape)
        if font_ref is not None and _first_color(font_ref) is not None:
            return self.theme.resolve(_first_color(font_ref), self.mapping)
        for lvl in default:
            solid = _def_rpr(lvl).find(p_qn("a:solidFill")) if _def_rpr(lvl) is not None else None
            if solid is not None:
                return self.theme.resolve(_first_color(solid), self.mapping)
        return None

    def run_attribute(self, shape, text_frame, level: int, name: str) -> str | None:
        """Atributo de defRPr (i, u, cap) heredado; None si nadie lo define."""
        main, default = self.levels(shape, text_frame, level)
        for lvl in main + default:
            def_rpr = _def_rpr(lvl)
            if def_rpr is not None and def_rpr.get(name) is not None:
                return def_rpr.get(name)
        return None

    def font(self, shape, text_frame, level: int) -> str | None:
        main, default = self.levels(shape, text_frame, level)
        for lvl in main:
            latin = _def_rpr(lvl).find(p_qn("a:latin")) if _def_rpr(lvl) is not None else None
            if latin is not None and latin.get("typeface"):
                return self.theme.font(latin.get("typeface"))
        font_ref = self._font_ref(shape)
        if font_ref is not None and font_ref.get("idx") in ("major", "minor"):
            return self.theme.fonts.get(font_ref.get("idx"))
        for lvl in default:
            latin = _def_rpr(lvl).find(p_qn("a:latin")) if _def_rpr(lvl) is not None else None
            if latin is not None and latin.get("typeface"):
                return self.theme.font(latin.get("typeface"))
        return self.theme.fonts.get("major" if _placeholder_type(shape) in _TITLE_PLACEHOLDERS else "minor")

    def paragraph_attribute(self, shape, text_frame, level: int, name: str) -> str | None:
        """Atributo del nivel (algn) heredado."""
        main, default = self.levels(shape, text_frame, level)
        for lvl in main + default:
            if lvl.get(name) is not None:
                return lvl.get(name)
        return None

    def line_spacing(self, shape, text_frame, level: int):
        """Elemento a:lnSpc heredado, o None."""
        main, default = self.levels(shape, text_frame, level)
        for lvl in main + default:
            ln_spc = lvl.find(p_qn("a:lnSpc"))
            if ln_spc is not None:
                return ln_spc
        return None


_JUSTIFIED_ALGN = {"just", "dist", "justLow", "thaiDist"}


def _pptx_spacing_value(ln_spc, font_size: float | None) -> float | None:
    """a:lnSpc a múltiplo de sencillo: spcPct en milésimas de % o spcPts en centésimas de punto."""
    if ln_spc is None:
        return 1.0
    pct = ln_spc.find(p_qn("a:spcPct"))
    if pct is not None and pct.get("val", "").isdigit():
        return int(pct.get("val")) / 100_000
    pts = ln_spc.find(p_qn("a:spcPts"))
    if pts is not None and pts.get("val", "").isdigit() and font_size:
        return int(pts.get("val")) / 100 / font_size
    return None


def _pptx_paragraph_format(shape, text_frame, paragraph, defaults: _PptxTextDefaults,
                           font_size: float | None) -> ParagraphFormat:
    p_pr = paragraph._p.pPr
    algn = p_pr.get("algn") if p_pr is not None else None
    if algn is None:
        algn = defaults.paragraph_attribute(shape, text_frame, paragraph.level, "algn")
    ln_spc = p_pr.find(p_qn("a:lnSpc")) if p_pr is not None else None
    if ln_spc is None:
        ln_spc = defaults.line_spacing(shape, text_frame, paragraph.level)
    spacing = _pptx_spacing_value(ln_spc, font_size)
    # "Reducir el texto al desbordar" también achica el interlineado (lnSpcReduction).
    autofit = text_frame._bodyPr.find(p_qn("a:normAutofit"))
    reduction = autofit.get("lnSpcReduction") if autofit is not None else None
    if spacing is not None and reduction and reduction.isdigit():
        spacing *= 1 - int(reduction) / 100_000
    return ParagraphFormat(justified=algn in _JUSTIFIED_ALGN, line_spacing=spacing)


def _pptx_run_attribute(run, paragraph, defaults, shape, text_frame, name: str) -> str | None:
    for props in (run._r.rPr, _def_rpr(paragraph._p.pPr)):
        if props is not None and props.get(name) is not None:
            return props.get(name)
    return defaults.run_attribute(shape, text_frame, paragraph.level, name)


def _pptx_run_font(run, paragraph, defaults, shape, text_frame, theme: _PptxTheme) -> str | None:
    for props in (run._r.rPr, _def_rpr(paragraph._p.pPr)):
        latin = props.find(p_qn("a:latin")) if props is not None else None
        if latin is not None and latin.get("typeface"):
            return theme.font(latin.get("typeface"))
    return defaults.font(shape, text_frame, paragraph.level)


def _run_color(run, paragraph, theme: _PptxTheme, mapping: dict, default: RGB | None) -> RGB | None:
    """Color del run, del párrafo, de enlace o el heredado. None si no se puede resolver."""
    for props in (run._r.rPr, paragraph._p.pPr.find(p_qn("a:defRPr")) if paragraph._p.pPr is not None else None):
        if props is None:
            continue
        solid = props.find(p_qn("a:solidFill"))
        if solid is not None:
            return theme.resolve(_first_color(solid), mapping)
        if any(props.find(p_qn(t)) is not None for t in ("a:gradFill", "a:blipFill", "a:pattFill")):
            return None
    r_pr = run._r.rPr
    if r_pr is not None and r_pr.find(p_qn("a:hlinkClick")) is not None:
        return theme.resolve(etree.fromstring(f'<a:schemeClr xmlns:a="{_A_NS["a"]}" val="hlink"/>'), mapping)
    return default


def _pptx_text_frames(shapes, theme: _PptxTheme, mapping: dict, slide_background: RGB | None):
    """(forma, marco de texto, fondo) de formas, grupos y celdas de tabla."""
    for shape in shapes:
        if isinstance(shape, GroupShape):
            yield from _pptx_text_frames(shape.shapes, theme, mapping, slide_background)
        elif getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells:
                    # Sin relleno propio, el de la celda sale del estilo de tabla: no se evalúa.
                    defined, fill = theme.fill(cell._tc.tcPr, mapping)
                    yield shape, cell.text_frame, fill if defined else None
        elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            defined, fill = _shape_fill(shape, theme, mapping)
            yield shape, shape.text_frame, fill if defined else slide_background


@_never_fails
def extract_pptx_legibility(presentation) -> LegibilityStats:
    """presentation es un pptx.Presentation ya abierto. Se saltean las diapositivas ocultas."""
    collector = _Collector()
    themes: dict[int, _PptxTheme] = {}
    for number, slide in enumerate(presentation.slides, start=1):
        if slide._element.get("show") in ("0", "false"):
            continue
        layout = slide.slide_layout
        master = layout.slide_master
        theme = themes.setdefault(id(master), _PptxTheme(master))
        mapping = _PptxTheme.color_map(theme.master_map, layout, slide)
        defaults = _PptxTextDefaults(presentation, theme, master, mapping)
        background = _pptx_slide_background(slide, theme, mapping)
        for shape, text_frame, fill in _pptx_text_frames(slide.shapes, theme, mapping, background):
            scale = _pptx_font_scale(text_frame)
            for paragraph in text_frame.paragraphs:
                default_color = defaults.color(shape, text_frame, paragraph.level)
                fragments = []
                for run in paragraph.runs:
                    size = run.font.size or paragraph.font.size
                    bold = run.font.bold if run.font.bold is not None else bool(paragraph.font.bold)
                    attribute = functools.partial(_pptx_run_attribute, run, paragraph, defaults, shape, text_frame)
                    r_pr = run._r.rPr
                    fragments.append(Fragment(
                        text=run.text,
                        size=size.pt * scale if size is not None else None,
                        color=_run_color(run, paragraph, theme, mapping, default_color),
                        background=fill,
                        bold=bold,
                        italic=attribute("i") in ("1", "true"),
                        underline=attribute("u") not in (None, "none"),
                        caps=attribute("cap") == "all",
                        font=_pptx_run_font(run, paragraph, defaults, shape, text_frame, theme),
                        link=r_pr is not None and r_pr.find(p_qn("a:hlinkClick")) is not None,
                    ))
                fmt = _pptx_paragraph_format(shape, text_frame, paragraph, defaults, _predominant_size(fragments))
                collector.add_paragraph(fragments, number, fmt)
    return collector.result()
